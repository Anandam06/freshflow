"""S21 Processor rescue unavailable — paired Excel-dataset uploads to the LIVE deployment.

Control and treatment workbooks are identical except Buyer_Capacity of every processor-channel buyer = 0.
"""
import base64
import io
import json
from datetime import datetime, timezone
from pathlib import Path

import httpx
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'evidence' / 'raw'
BASE = 'https://freshflow-tatz.onrender.com'

BATCHES = [
    ('B-001', 'tomato', 'H1', 'B02', 800, 'A', [('refrigeration_failure', 'high'), ('temperature_rise', 'high')]),
    ('B-002', 'spinach', 'H1', 'B05', 400, 'A', [('temperature_rise', 'high')]),
    ('B-003', 'tomato', 'H1', 'B02', 800, 'C', [('product_damage', 'high')]),
    ('B-004', 'mango', 'H2', 'B09', 500, 'A', [('buyer_cancellation', 'high')]),
    ('B-005', 'brinjal', 'H1', 'B05', 400, 'B', [('refrigeration_failure', 'high')]),
    ('B-006', 'tomato', 'H1', 'B02', 800, 'B', [('buyer_cancellation', 'high'), ('heavy_traffic', 'high')]),
]


def workbook(template, processors_off):
    wb = openpyxl.load_workbook(io.BytesIO(template))
    inv = wb['Batch_Inventory']
    inv.delete_rows(2, inv.max_row)
    for b in BATCHES:
        inv.append(list(b[:6]))
    ev = wb['Disruption_Events']
    ev.delete_rows(2, ev.max_row)
    for b in BATCHES:
        for t, s in b[6]:
            ev.append([b[0], t, s])
    channel = {r[0]: r[3] for r in wb['Locations'].iter_rows(min_row=2, values_only=True)}
    zeroed = []
    if processors_off:
        for row in wb['Buyer_Capacity'].iter_rows(min_row=2):
            if channel.get(row[0].value) == 'processor':
                row[1].value = 0
                zeroed.append(row[0].value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), zeroed, channel


def main():
    c = httpx.Client(base_url=BASE, timeout=300)
    template = c.get('/api/v1/excel/template').content
    result = {'executed_utc': datetime.now(timezone.utc).isoformat(timespec='seconds'), 'arms': {}}
    for arm, off in [('control', False), ('processors_unavailable', True)]:
        content, zeroed, channel = workbook(template, off)
        (OUT / f's21_{arm}.xlsx').write_bytes(content)
        r = c.post('/api/v1/excel/upload', json={'filename': f's21_{arm}.xlsx', 'content_base64': base64.b64encode(content).decode()})
        body = r.json()
        full = [c.get(f'/api/v1/shipments/{s["shipment_id"]}').json() for s in body.get('shipments', [])] if r.status_code == 200 else []
        result['arms'][arm] = {'http_status': r.status_code, 'zeroed_processors': zeroed, 'run': body, 'shipments': full}
        result['channels'] = channel
        print(arm, r.status_code, body.get('id'), [s.get('decision_status') for s in body.get('shipments', [])])
    (OUT / 's21_live.json').write_text(json.dumps(result, indent=1), encoding='utf-8')


if __name__ == '__main__':
    main()
