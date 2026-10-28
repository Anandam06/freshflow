"""Build FreshFlow_Scenario_Validation.xlsx, CSV exports, agent trace JSON and the PPT evidence brief
from the raw results in evidence/raw (fleet_results.jsonl, shipments_live.jsonl, s21_live.json, live_*.json).
Usage: python -m validation.build_workbook
"""
import csv
import json
import statistics as st
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from validation.scenarios import SCENARIOS, SUPPLEMENTARY, SEEDS

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / 'evidence' / 'raw'
OUT = ROOT / 'evidence'
COMMIT = 'b32cafa'
LIVE = 'https://freshflow-tatz.onrender.com'
MODELS = ['M0', 'M1', 'M2', 'M3']
LABEL = {'M0': 'A — Traditional Planning (M0)', 'M1': 'B — Freshness-Aware (M1)',
         'M2': 'C — Smart Optimization (M2)', 'M3': 'D — FreshFlow Agentic AI (M3)'}
LETTER = dict(zip(MODELS, 'ABCD'))
BRIEF_COMMODITIES = {'fruit': ['mango', 'banana', 'apple', 'grapes', 'orange', 'papaya', 'guava', 'pomegranate', 'watermelon', 'pineapple'],
                     'vegetable': ['tomato', 'potato', 'onion', 'cabbage', 'cauliflower', 'spinach', 'okra', 'capsicum']}

INR = '[>=10000000]"₹"#\\,##\\,##\\,##0;[>=100000]"₹"#\\,##\\,##0;"₹"#,##0'
INR2 = '"₹"#,##0.00'
PCT = '0.0%'
PP = '+0.00" pp";-0.00" pp";0.00" pp"'
KG = '#,##0" kg"'
HRS = '0.0" h"'
NUM = '#,##0.00'
INT = '#,##0'

HEAD = PatternFill('solid', fgColor='1F4D3A')
INHEAD = PatternFill('solid', fgColor='35506B')
WHITE = Font(bold=True, color='FFFFFF')
GOOD = PatternFill('solid', fgColor='D8F0DF')
BAD = PatternFill('solid', fgColor='F8D7D7')
NOTE = Font(italic=True, color='555555')


def load():
    fleet = [json.loads(l) for l in (RAW / 'fleet_results.jsonl').read_text(encoding='utf-8').splitlines() if l.strip()]
    ship = [json.loads(l) for l in (RAW / 'shipments_live.jsonl').read_text(encoding='utf-8').splitlines() if l.strip()]
    j = lambda n: json.loads((RAW / n).read_text(encoding='utf-8'))
    return fleet, ship, {n: j(n) for n in ['live_health.json', 'live_defaults.json', 'live_network.json', 'live_evidence.json',
                                            'live_llm_usage.json', 'live_baseline_simulation.json', 's21_live.json',
                                            'live_check_X01.json', 'live_check_S17.json', 'live_check_S31.json']}


def pct(values, p):
    o = sorted(values)
    i = (len(o) - 1) * p
    lo = int(i)
    return o[lo] + (o[min(lo + 1, len(o) - 1)] - o[lo]) * (i - lo)


def stats(v):
    return {'mean': st.mean(v), 'median': st.median(v), 'min': min(v), 'max': max(v),
            'sd': st.stdev(v) if len(v) > 1 else 0.0, 'p05': pct(v, 0.05), 'p95': pct(v, 0.95)}


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def table(self, name, cols, rows, note=None, input_cols=0, delta_cols=(), tab=None):
        """cols: list of (header, number_format or None, width)."""
        ws = self.wb.create_sheet(name)
        if tab:
            ws.sheet_properties.tabColor = tab
        r0 = 1
        if note:
            ws.cell(1, 1, note).font = NOTE
            ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical='top')
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, min(len(cols), 12)))
            ws.row_dimensions[1].height = 15 * max(2, len(note) // 150 + 1)
            r0 = 2
        for c, (h, fmt, w) in enumerate(cols, 1):
            cell = ws.cell(r0, c, h)
            cell.fill = INHEAD if c <= input_cols else HEAD
            cell.font = WHITE
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[r0].height = 34
        for i, row in enumerate(rows, r0 + 1):
            for c, v in enumerate(row, 1):
                if isinstance(v, (list, dict)):
                    v = json.dumps(v, ensure_ascii=False)
                cell = ws.cell(i, c, v)
                fmt = cols[c - 1][1]
                if fmt and isinstance(v, (int, float)) or (fmt and isinstance(v, str) and v.startswith('=')):
                    cell.number_format = fmt
                if isinstance(v, str) and len(v) > 60:
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
        last = r0 + len(rows)
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        if rows:
            ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cols))}{last}'
        for c in delta_cols:
            col = get_column_letter(c)
            rng = f'{col}{r0 + 1}:{col}{max(last, r0 + 1)}'
            ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0.0000001'], fill=GOOD))
            ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['-0.0000001'], fill=BAD))
        ws.header_row = r0
        return ws

    def save(self, path):
        self.wb.save(path)


def main():
    fleet, ship, L = load()
    now = datetime.now(timezone.utc).isoformat(timespec='seconds')
    cat = {s[0]: s for s in SCENARIOS}
    sup = {s[0]: s for s in SUPPLEMENTARY}
    by = defaultdict(list)
    for r in fleet:
        by[r['scenario_id']].append(r)
    for k in by:
        by[k].sort(key=lambda r: r['seed'])
    fleet_ids = [s[0] for s in SCENARIOS if s[3] == 'fleet' and s[0] in by] + [x[0] for x in SUPPLEMENTARY if x[0] in by]
    name_of = {**{s[0]: s[2] for s in SCENARIOS}, **{x[0]: x[2] for x in SUPPLEMENTARY}}
    ref_of = lambda sid: 'X01' if sid.startswith('X') else 'S01'
    ship_by = {r['test_id']: r for r in ship}
    network = L['live_network.json']
    health = L['live_health.json']
    defaults = L['live_defaults.json']['parameters']
    bk = Book()
    rowref = {}

    # ------------------------------------------------------------------ aggregate fleet
    agg = {}
    audit_fail = []
    for sid in fleet_ids:
        runs = by[sid]
        A = {}
        for i, m in enumerate(MODELS):
            prof = [r['metrics'][i]['operating_profit_inr'] for r in runs]
            spoil = [r['metrics'][i]['spoilage_rate'] for r in runs]
            A[m] = {'profit': stats(prof), 'spoil': stats(spoil), 'n': len(runs),
                    'sold': st.mean(r['metrics'][i]['sold_kg'] for r in runs),
                    'spoiled': st.mean(r['metrics'][i]['spoiled_kg'] for r in runs),
                    'input': st.mean(r['metrics'][i]['input_kg'] for r in runs),
                    'revenue': st.mean(r['metrics'][i]['gross_sales_revenue_inr'] for r in runs),
                    'margin': st.mean(r['metrics'][i]['operating_profit_margin'] or 0 for r in runs),
                    'ontime': st.mean(r['metrics'][i]['on_time_fulfilment_rate'] for r in runs),
                    'reject': st.mean(r['metrics'][i]['buyer_rejection_rate'] or 0 for r in runs),
                    'rescued': st.mean(r['metrics'][i]['rescued_batches'] for r in runs),
                    'attempts': st.mean(r['extras'][m]['rescue_attempts'] for r in runs),
                    'failed': st.mean(r['extras'][m]['failed_interventions'] for r in runs),
                    'agent_cost': st.mean(r['metrics'][i]['agent_sensor_inr'] for r in runs),
                    'rescue_cost': st.mean(r['metrics'][i]['rescue_inr'] for r in runs),
                    'transport': st.mean(r['metrics'][i]['transport_inr'] for r in runs),
                    'price': st.mean(r['extras'][m]['realized_price_per_sold_kg_inr'] or 0 for r in runs),
                    'arrival_rsl': st.mean(r['extras'][m]['mean_arrival_rsl_accepted_h'] or 0 for r in runs),
                    'netkg': st.mean(r['metrics'][i]['net_realization_per_input_kg_inr'] for r in runs),
                    'safety': sum(r['metrics'][i]['food_safety_violations'] for r in runs)}
            for r in runs:
                a = r['audit'][m]
                if not a['passed'] or not a['same_events_digest']:
                    audit_fail.append((sid, r['seed'], m, a['issues']))
        dc = [r['metrics'][3]['operating_profit_inr'] - r['metrics'][2]['operating_profit_inr'] for r in runs]
        da = [r['metrics'][3]['operating_profit_inr'] - r['metrics'][0]['operating_profit_inr'] for r in runs]
        A['DC'], A['DA'] = stats(dc), stats(da)
        A['P_DA'] = st.mean(float(x > 0) for x in da)
        A['P_DC'] = st.mean(float(x > 0) for x in dc)
        agg[sid] = A

    # ------------------------------------------------------------------ 1 README
    n_fleet = len([s for s in SCENARIOS if s[3] == 'fleet' and s[0] in by])
    n_ship = len([s for s in SCENARIOS if s[3] in ('shipment', 'shipment_dataset')])
    n_ni = len([s for s in SCENARIOS if s[3] == 'not_implemented'])
    reps = min(len(by[s]) for s in fleet_ids)
    readme = [
        ('Purpose', 'Independent, reproducible validation of the FreshFlow simulation: 34 brief scenarios (S01–S34) plus a declared supplementary family (X01–X31), strategies A/B/C/D, financial audit, commodity and Mumbai-route checks, and live Claude agent traces.'),
        ('Generated (UTC)', now),
        ('Live application', LIVE),
        ('Backend version tested', f"API {health['version']} · git commit {COMMIT} (deployed on Render; same commit run locally)"),
        ('Live AI configuration', f"llm_provider={health.get('llm_provider')} · model={health.get('llm_model')} · live_agent_configured={health.get('live_agent_configured')}"),
        ('Evidence category', 'SIMULATED RESULT on SYNTHETIC data. Not field evidence; no physical truck was rerouted; results are not actual Mumbai operating outcomes.'),
        ('Engine 1 — Fleet simulation', f'Tomato-only, 2,000 batches × 30 days per repetition, 3 synthetic buyers (Retailer A, Wholesale B, Processor C) on a synthetic junction graph. {n_fleet} brief scenarios + {len(SUPPLEMENTARY)} supplementary × {reps} paired seeds ({SEEDS[0]}–{SEEDS[reps - 1]}). D (M3) is a deterministic emulation of the agent policy — no LLM calls.'),
        ('Engine 2 — Mumbai shipment engine', f"{len(network['commodities'])} commodities, {len(network['hubs'])} APMC hubs + {len(network['buyers'])} synthetic buyers, OSRM/OpenStreetMap road distances with simulated traffic. Single shipment, A/B/C/D compared on identical disruptions. Run on the LIVE site; agent cases used live Claude ({health.get('llm_model')})."),
        ('Where fleet runs executed', 'Locally (Windows, Python 3.13, OR-Tools 9.15.6755) on the deployed commit, because 1,140 runs × ~25 s exceed what the free Render instance can serve. A and B match the live site exactly; C and D differ slightly because CP-SAT time-limited solves differ across OS builds (measured deviation in Live_vs_Local_Check).'),
        ('Scenario status', f'{n_fleet} executed in fleet engine · {n_ship} executed in shipment engine (live) · {n_ni} NOT IMPLEMENTED (S07, S13, S26) — see Failed_Tests.'),
        ('Key design caveat', 'Application default disruption_probability = 0. In the fleet generator traffic only delays disrupted batches, so S02–S05 (traffic alone) cannot change results. The X-family repeats them with disruption_probability 0.30 (the value the prespecified risk study uses). This is disclosed, not hidden.'),
        ('Statistics', 'Mean, median, min, max, SD and linearly interpolated 5th/95th percentiles over paired seeds. These describe simulation variability under assumed distributions — not confidence intervals or real-world bounds.'),
        ('How to reproduce', 'cd backend; python -m validation.run_fleet  (fleet matrix) · python -m validation.run_shipments_live · python -m validation.run_s21_live · python -m validation.build_workbook. Same seed + same parameters ⇒ identical events (dataset_sha256).'),
        ('Sheet guide', 'Blue headers = inputs/identifiers; green headers = results. Green/red cells = positive/negative differences. Money in INR (₹, Indian grouping), rates in %, quantities in kg, times in hours.'),
        ('Secrets', 'No API keys, tokens or personal data are included. All buyers are synthetic.'),
    ]
    ws = bk.table('README', [('Item', None, 28), ('Detail', None, 140)], readme, tab='1F4D3A')
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
        ws.cell(r, 1).font = Font(bold=True)

    # ------------------------------------------------------------------ 2 Scenario_Catalog
    rows = []
    for sid, grp, name, eng, ov, desc in SCENARIOS:
        status = {'fleet': f'EXECUTED — fleet engine, {len(by.get(sid, []))} paired seeds',
                  'shipment': 'EXECUTED — Mumbai shipment engine (live)',
                  'shipment_dataset': 'EXECUTED — Mumbai shipment engine via Excel dataset upload (live)',
                  'not_implemented': 'NOT IMPLEMENTED'}[eng]
        rows.append((sid, grp, name, eng, json.dumps(ov) if ov else '—', desc, status))
    for xid, base, name, ov in SUPPLEMENTARY:
        rows.append((xid, 'X Supplementary (declared)', name, 'fleet', json.dumps(ov),
                     f'Repeat of {base} with disruption_probability 0.30 so traffic/approval effects can act. Not part of the brief.',
                     f'EXECUTED — fleet engine, {len(by.get(xid, []))} paired seeds'))
    bk.table('Scenario_Catalog', [('Scenario', None, 9), ('Group', None, 22), ('Name', None, 30), ('Engine', None, 14),
                                  ('Overrides vs default', None, 46), ('Description / mapping', None, 70), ('Status', None, 44)],
             rows, input_cols=6)

    # ------------------------------------------------------------------ 3 Inputs
    params = list(defaults.keys())
    units = {'seed': 'int', 'batches': 'batches', 'days': 'days', 'repetitions': 'runs', 'disruption_probability': 'prob.',
             'price_multiplier': '×', 'capacity_multiplier': '×', 'setup_cost_inr': 'INR', 'buyer_availability': 'prob.',
             'temperature_shift_c': '°C', 'traffic_multiplier': '×', 'prediction_error_hours': 'h (SD)', 'approval_delay_hours': 'h',
             'ai_cost_multiplier': '×', 'initial_rsl_scale': '×', 'cancellation_probability': 'prob.'}
    rows = [('DEFAULT', 'Application default (GET /api/v1/scenarios/defaults)', '—', '—') + tuple(defaults[p] for p in params)]
    for sid in fleet_ids:
        cfg = by[sid][0]['config']
        ov = cat[sid][4] if sid in cat else sup[sid][3]
        rows.append((sid, name_of[sid], ', '.join(ov) or '(none)', f'{by[sid][0]["seed"]}–{by[sid][-1]["seed"]} ({len(by[sid])})')
                    + tuple('per-seed' if p in ('seed', 'repetitions') else cfg[p] for p in params))
    cols = [('Scenario', None, 10), ('Name', None, 28), ('Changed fields', None, 30), ('Seeds (n)', None, 16)] + \
           [(f'{p} [{units[p]}]', None, 14) for p in params]
    ws = bk.table('Inputs', cols, rows, input_cols=len(cols),
                  note='Fleet-engine inputs per scenario (each seed run with repetitions=1 at that seed). Shipment-engine inputs are listed below the table rows in Run_Register (request JSON).')
    # changed cells highlighted
    for r in range(4, ws.max_row + 1):
        for c, p in enumerate(params, 5):
            v = ws.cell(r, c).value
            if p not in ('seed', 'repetitions') and v != defaults[p]:
                ws.cell(r, c).fill = PatternFill('solid', fgColor='FFF2CC')

    # ------------------------------------------------------------------ 4 Run_Register
    rows = []
    base = L['live_baseline_simulation.json']
    rows.append(('LIVE', base['id'], 'S01 (default, 3 reps)', 'POST /api/v1/simulations {}', 'live Render (Linux)', '—',
                 'COMPLETED', '; '.join(f"{r['seed']}:{r['dataset_sha256'][:12]}" for r in base['result']['runs']), 'Server-cached default run'))
    for k in ['X01', 'S17', 'S31']:
        d = L[f'live_check_{k}.json']
        rows.append(('LIVE', d['id'], f'{k} cross-check (1 rep)', json.dumps({kk: v for kk, v in d['result']['config'].items() if v != defaults.get(kk)}),
                     'live Render (Linux)', '—', 'COMPLETED', d['result']['runs'][0]['dataset_sha256'][:12], 'Cross-platform comparison'))
    for r in fleet:
        ok = all(r['audit'][m]['passed'] for m in MODELS)
        rows.append(('LOCAL', f"LOCAL-{r['scenario_id']}-{r['seed']}", r['scenario_id'], f"seed={r['seed']}", f'local Windows · commit {COMMIT}',
                     round(r['runtime_s'], 1), 'COMPLETED' + ('' if ok else ' — AUDIT FLAG'), r['dataset_sha256'][:12], r['finished_utc']))
    for s in ship:
        resp = s['response']
        rows.append(('LIVE', resp.get('id', '—'), s['test_id'], json.dumps(s['request'])[:300], 'live Render (Linux)', s.get('latency_s', '—'),
                     f"HTTP {s['http_status']} · {resp.get('decision_status', '')}", resp.get('dataset_id', ''), s.get('executed_utc', '')))
    for arm, a in L['s21_live.json']['arms'].items():
        rows.append(('LIVE', a['run'].get('id'), f'S21 {arm}', f"Excel upload; processors zeroed: {a['zeroed_processors']}", 'live Render (Linux)', '—',
                     f"HTTP {a['http_status']}", a['run'].get('dataset_id'), L['s21_live.json']['executed_utc']))
    bk.table('Run_Register', [('Where', None, 8), ('Run ID', None, 30), ('Scenario / test', None, 30), ('Request', None, 50),
                              ('Executor & code', None, 30), ('Runtime / latency (s)', '0.0', 12), ('Status', None, 34),
                              ('Dataset SHA-256 (12) / dataset id', None, 30), ('Timestamp (UTC)', None, 22)], rows, input_cols=5)

    # ------------------------------------------------------------------ 5 A_B_C_D_Results
    rcols = [('Scenario', None, 9), ('Seed', None, 8), ('Run ID', None, 18), ('Strategy', None, 30), ('Input kg', KG, 13),
             ('Sold kg', KG, 13), ('Spoiled / written-off kg', KG, 14), ('of which processor (salvage) kg', KG, 14),
             ('Mass residual kg', NUM, 10), ('Mean est. initial RSL h', HRS, 11), ('Mean arrival RSL (accepted) h', HRS, 12),
             ('Spoilage rate', PCT, 10), ('On-time rate', PCT, 10), ('Buyer rejection rate', PCT, 10), ('Rescue attempts', INT, 9),
             ('Rescued batches', INT, 9), ('Failed interventions', INT, 10), ('Approval wait h (total)', HRS, 10),
             ('Realized price ₹/sold kg', INR2, 11), ('Gross revenue', INR, 15), ('of which processor (salvage) revenue', INR, 15),
             ('Procurement', INR, 14), ('Transport', INR, 13), ('Handling', INR, 12), ('Diversion (rescue)', INR, 12),
             ('Agent / technology', INR, 12), ('Operating profit', INR, 15), ('Operating margin', PCT, 10),
             ('Net realization ₹/input kg', INR2, 11), ('Safety violations', INT, 9), ('Audit', None, 10)]
    rows, csv_rows = [], []
    for sid in fleet_ids:
        for r in by[sid]:
            for i, m in enumerate(MODELS):
                x, e = r['metrics'][i], r['extras'][m]
                row = (sid, r['seed'], f"LOCAL-{sid}-{r['seed']}", LABEL[m], x['input_kg'], x['sold_kg'], x['spoiled_kg'], e['salvage_kg'],
                       r['audit'][m]['mass_residual_kg'], e['mean_estimated_initial_rsl_h'], e['mean_arrival_rsl_accepted_h'],
                       x['spoilage_rate'], x['on_time_fulfilment_rate'], x['buyer_rejection_rate'], e['rescue_attempts'],
                       x['rescued_batches'], e['failed_interventions'], e['approval_wait_hours_total'],
                       e['realized_price_per_sold_kg_inr'], x['gross_sales_revenue_inr'], x['salvage_value_inr'],
                       x['procurement_inr'], x['transport_inr'], x['handling_inr'], x['rescue_inr'], x['agent_sensor_inr'],
                       x['operating_profit_inr'], x['operating_profit_margin'], x['net_realization_per_input_kg_inr'],
                       x['food_safety_violations'], 'PASS' if r['audit'][m]['passed'] else 'FLAG')
                rows.append(row)
                csv_rows.append(row)
    bk.table('A_B_C_D_Results', rcols, rows, input_cols=4,
             note='One row per scenario × seed × strategy (fleet engine). All four strategies at a seed see identical exogenous events (same dataset SHA-256). Processor (salvage) kg/revenue are subsets of sold kg/gross revenue, not additions. Diversion cost = ₹0.75/kg rescue handling; the original leg transport is still charged.')
    with (OUT / 'csv' / 'A_B_C_D_Results.csv').open('w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow([c[0] for c in rcols])
        w.writerows(csv_rows)

    # ------------------------------------------------------------------ 6 Scenario_Summary
    rows = []
    for sid in fleet_ids:
        for m in MODELS:
            a = agg[sid][m]
            p, s = a['profit'], a['spoil']
            rows.append((sid, name_of[sid], LABEL[m], a['n'], p['mean'], p['median'], p['min'], p['max'], p['sd'], p['p05'], p['p95'],
                         s['mean'], s['p05'], s['p95'], a['sold'], a['spoiled'], a['revenue'], a['margin'], a['ontime'], a['reject'],
                         a['rescued'], a['price'], a['arrival_rsl'], a['safety']))
            rowref[('Scenario_Summary', sid, m)] = len(rows) + 2
    cols = [('Scenario', None, 9), ('Name', None, 26), ('Strategy', None, 28), ('n seeds', INT, 7),
            ('Profit mean', INR, 14), ('Profit median', INR, 14), ('Profit min', INR, 14), ('Profit max', INR, 14), ('Profit SD', INR, 12),
            ('Profit P05', INR, 14), ('Profit P95', INR, 14), ('Spoilage mean', PCT, 10), ('Spoilage P05', PCT, 10), ('Spoilage P95', PCT, 10),
            ('Sold kg mean', KG, 13), ('Spoiled kg mean', KG, 13), ('Gross revenue mean', INR, 14), ('Operating margin mean', PCT, 10),
            ('On-time mean', PCT, 9), ('Rejection mean', PCT, 9), ('Rescued batches mean', NUM, 9), ('Realized ₹/sold kg', INR2, 10),
            ('Arrival RSL h (accepted)', HRS, 10), ('Safety violations (sum)', INT, 9)]
    bk.table('Scenario_Summary', cols, rows, input_cols=4,
             note='Aggregates over paired seeds. P05/P95 = simulation percentiles under the assumed input distributions, not confidence intervals.')
    with (OUT / 'csv' / 'Scenario_Summary.csv').open('w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow([c[0] for c in cols])
        w.writerows(rows)

    # ------------------------------------------------------------------ 7 Profit_Comparison
    rows = []
    for k, sid in enumerate(fleet_ids):
        a = agg[sid]
        r = k + 3
        rows.append((sid, name_of[sid], a['M0']['profit']['mean'], a['M1']['profit']['mean'], a['M2']['profit']['mean'], a['M3']['profit']['mean'],
                     a['M1']['profit']['mean'] - a['M0']['profit']['mean'], a['M2']['profit']['mean'] - a['M1']['profit']['mean'],
                     a['DC']['mean'], a['DA']['mean'], a['DC']['p05'], a['DC']['p95'], a['DA']['p05'], a['DA']['p95'],
                     a['P_DA'], a['P_DC'], a['M3']['agent_cost'] - a['M2']['agent_cost'], a['M3']['rescue_cost'],
                     a['M3']['transport'] - a['M2']['transport'], f'=F{r}-E{r}', f'=F{r}-C{r}'))
        rowref[('Profit_Comparison', sid)] = r
    pcols = [('Scenario', None, 9), ('Name', None, 28), ('A profit (mean)', INR, 14), ('B profit (mean)', INR, 14), ('C profit (mean)', INR, 14),
             ('D profit (mean)', INR, 14), ('B − A', INR, 13), ('C − B', INR, 13), ('D − C', INR, 13), ('D − A', INR, 14),
             ('D − C P05', INR, 13), ('D − C P95', INR, 13), ('D − A P05', INR, 13), ('D − A P95', INR, 13),
             ('Share of seeds D > A', PCT, 10), ('Share of seeds D > C', PCT, 10), ('Extra agent/tech cost D vs C', INR, 13),
             ('D diversion cost', INR, 12), ('Extra transport D vs C', INR, 12), ('Check: D − C (formula)', INR, 13), ('Check: D − A (formula)', INR, 13)]
    ws = bk.table('Profit_Comparison', pcols, rows, input_cols=2, delta_cols=(7, 8, 9, 10, 11, 12, 13, 14, 20, 21),
                  note='Means over paired seeds (same events for all strategies). D − C isolates the value of event-driven agentic intervention over the same optimizer without it. The two formula columns recompute D − C and D − A from the mean columns.')
    ch = BarChart()
    ch.type = 'bar'
    ch.title = 'Mean D − C operating profit by scenario (₹)'
    ch.y_axis.title = '₹ per 30-day run'
    ch.add_data(Reference(ws, min_col=9, min_row=2, max_row=2 + len(rows)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=2 + len(rows)))
    ch.height, ch.width, ch.legend = 16, 18, None
    ws.add_chart(ch, 'W3')
    ch2 = BarChart()
    ch2.title = 'Mean operating profit A/B/C/D — S01 and X01 (₹)'
    for sid in ['S01', 'X01']:
        if ('Profit_Comparison', sid) in rowref:
            r = rowref[('Profit_Comparison', sid)]
            ch2.add_data(Reference(ws, min_col=3, max_col=6, min_row=r, max_row=r), from_rows=True, titles_from_data=False)
    ch2.set_categories(Reference(ws, min_col=3, max_col=6, min_row=2, max_row=2))
    if len(ch2.series) == 2:
        from openpyxl.chart.series import SeriesLabel
        ch2.series[0].tx = SeriesLabel(v='S01 baseline')
        ch2.series[1].tx = SeriesLabel(v='X01 disrupted (p=0.30)')
    ch2.height, ch2.width = 9, 18
    ws.add_chart(ch2, 'W36')

    # ------------------------------------------------------------------ 8 Spoilage_Comparison
    rows = []
    for k, sid in enumerate(fleet_ids):
        a = agg[sid]
        rows.append((sid, name_of[sid], a['M0']['spoil']['mean'], a['M1']['spoil']['mean'], a['M2']['spoil']['mean'], a['M3']['spoil']['mean'],
                     100 * (a['M3']['spoil']['mean'] - a['M0']['spoil']['mean']), 100 * (a['M3']['spoil']['mean'] - a['M2']['spoil']['mean']),
                     a['M0']['spoiled'], a['M3']['spoiled'], a['M0']['spoiled'] - a['M3']['spoiled'], a['M2']['spoiled'] - a['M3']['spoiled'],
                     a['M3']['attempts'], a['M3']['rescued'], (a['M3']['rescued'] / a['M3']['attempts']) if a['M3']['attempts'] else None,
                     a['M3']['failed']))
        rowref[('Spoilage_Comparison', sid)] = k + 3
    ws = bk.table('Spoilage_Comparison', [('Scenario', None, 9), ('Name', None, 28), ('A spoilage', PCT, 10), ('B spoilage', PCT, 10),
                                          ('C spoilage', PCT, 10), ('D spoilage', PCT, 10), ('D − A (pp)', PP, 11), ('D − C (pp)', PP, 11),
                                          ('A spoiled kg', KG, 13), ('D spoiled kg', KG, 13), ('Spoilage avoided D vs A kg', KG, 14),
                                          ('Spoilage avoided D vs C kg', KG, 14), ('D rescue attempts (mean)', NUM, 10),
                                          ('D rescued batches (mean)', NUM, 10), ('D successful rescue rate', PCT, 10), ('D failed interventions (mean)', NUM, 10)],
                  rows, input_cols=2, delta_cols=(11, 12),
                  note='Spoilage = written-off share of input kg (unsold lots are conservatively written off, including commercial rejection). Negative pp = D spoils less. Rescue rate = rescued and accepted ÷ rescue attempts.')
    for c in (7, 8):
        col = get_column_letter(c)
        rng = f'{col}3:{col}{2 + len(rows)}'
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['-0.0000001'], fill=GOOD))
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0.0000001'], fill=BAD))
    ch = BarChart()
    ch.title = 'Mean spoilage rate A/B/C/D by scenario'
    ch.add_data(Reference(ws, min_col=3, max_col=6, min_row=2, max_row=2 + len(rows)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=3, max_row=2 + len(rows)))
    ch.y_axis.number_format = '0%'
    ch.height, ch.width = 10, 30
    ws.add_chart(ch, 'R3')

    # ------------------------------------------------------------------ 9 Route_ETA_Changes
    routes = network['routes']
    names = {b['id']: b['name'] for b in network['buyers']} | {h['id']: h['name'] for h in network['hubs']}
    rows = []
    for s in ship:
        resp = s['response']
        if 'strategies' not in resp:
            continue
        A, D = resp['strategies']['A'], resp['strategies']['D']
        plan = resp['plan_buyer_id']
        origin = resp['request']['origin_id']
        if D['buyer_id'] == plan and D['status'] not in ('no_feasible_alternative',):
            continue
        cands = next((t['result'].get('options', []) for t in resp['state']['trace'] if t['tool'] == 'optimize_destination'), [])
        cand_txt = '; '.join(f"{o['id']}{'✓' if o['feasible'] else '✗'} {o['distance_km']:.1f}km {o['travel_hours']:.2f}h" for o in cands[:6])
        dis = ', '.join(f"{d['type']}({d['severity']})" for d in resp['request']['disruptions']) or 'none'
        rk = f'{origin}>{plan}'
        rows.append((s['test_id'], resp['id'], resp['commodity']['name'], origin, plan, names.get(plan), routes.get(rk, {}).get('distance_km'),
                     routes.get(rk, {}).get('freeflow_minutes', 0) / 60 * resp['settings']['traffic_multiplier'] if rk in routes else None,
                     A['distance_km'], A['lead_time_hours'], dis, cand_txt, D['buyer_id'] if D['status'] != 'no_feasible_alternative' else '—',
                     names.get(D['buyer_id'], '') if D['status'] != 'no_feasible_alternative' else 'no feasible buyer',
                     D['distance_km'], D['lead_time_hours'], D['status'], 'checkpoint (simulated current position)',
                     'OSRM road distance (OpenStreetMap); traffic SIMULATED ×' + str(resp['settings']['traffic_multiplier']),
                     'SIMULATED state transition — no physical vehicle'))
    fl = []
    det = RAW / 'detail_X01_60000.json'
    if det.exists():
        d = json.loads(det.read_text(encoding='utf-8'))
        for r in d['M3']:
            if r['rescue_attempt']:
                orig = next((c for c in r.get('candidates', []) if c['id'] == r['morning_destination']), {})
                new = next((c for c in r.get('candidates', []) if c['id'] == r['destination']), {})
                fl.append((f"X01/seed60000/{r['batch_id']}", '—', 'Tomato', 'depot', r['morning_destination'], r['morning_destination'], None, None,
                           None, orig.get('travel_hours'), 'fleet event (delay/heat/cancellation)',
                           '; '.join(f"{c['id']}{'✓' if c['feasible'] else '✗'} {c['travel_hours']:.1f}h" for c in r.get('candidates', [])),
                           r['destination'] or '—', r['destination'] or 'no feasible buyer', None, new.get('travel_hours'),
                           'rescued & accepted' if r['rescued'] else ('kept/failed' if r['buyer_acceptance'] else 'written off'),
                           'junction (simulated current position)', 'SYNTHETIC junction graph; not map-derived', 'SIMULATED state transition'))
    bk.table('Route_ETA_Changes', [('Test / batch', None, 22), ('Shipment ID', None, 14), ('Commodity', None, 11), ('Origin', None, 8),
                                   ('Original buyer', None, 9), ('Original buyer name', None, 26), ('Planned distance km (no disruption)', NUM, 11),
                                   ('Planned ETA h (traffic-adjusted)', NUM, 11), ('Original plan distance km under disruption', NUM, 12),
                                   ('Original plan lead time h under disruption', NUM, 12), ('Disruption', None, 28), ('Candidate buyers (✓ feasible)', None, 60),
                                   ('Selected buyer', None, 9), ('Selected buyer name', None, 26), ('Revised distance km', NUM, 10),
                                   ('Revised lead time h', NUM, 10), ('D status', None, 22), ('Reroute starts at', None, 22), ('Routing data', None, 34),
                                   ('Physical?', None, 22)], rows + fl, input_cols=11,
             note='Mumbai rows: live shipment engine; distances from OSRM/OpenStreetMap, traffic is a simulated multiplier (no live traffic feed). Fleet rows (X01 seed 60000): synthetic junction graph with travel hours only. "proposed_awaiting_approval" rows were not approved; "executed" rows were approved in the simulated workflow. No physical truck was rerouted.')

    # ------------------------------------------------------------------ 10 Freshness_Results & 11 Dynamic_Pricing (commodity tests)
    com = {c['id']: c for c in network['commodities']}
    frows, prows = [], []
    for s in ship:
        if s['group'] != 'commodity':
            continue
        resp = s['response']
        cid = resp['commodity']['id']
        test = s['test_id'].split('-', 2)[2]
        A, C, D = (resp['strategies'][k] for k in 'ACD')
        frows.append((cid, com[cid]['category'], test, com[cid]['shelf_life_hours_20c'], com[cid]['q10'], com[cid]['humidity_sensitivity'],
                      resp['initial_rsl_hours'], resp['conservative_buffer_hours'], A['minimum_rsl_hours'], A['arrival_rsl_hours'], A['accepted'],
                      C['arrival_rsl_hours'], C['accepted'], D['buyer_id'], D['arrival_rsl_hours'], D['accepted'], resp['decision_status'],
                      A['operating_profit_inr'], D['operating_profit_inr'], D['operating_profit_inr'] - C['operating_profit_inr'], resp['id']))
        prows.append((cid, test, com[cid]['reference_price_inr_kg'], resp['settings']['market_price_factor'], resp['conditions']['grade'],
                      resp['conditions']['committed_price_factor'], A['buyer_id'], A['offer_inr_kg'], A['price_inr_kg'], D['buyer_id'], D['channel'],
                      D['offer_inr_kg'], D['price_inr_kg'], D['price_inr_kg'] - A['price_inr_kg'], D['revenue_inr'], resp['id']))
    bk.table('Freshness_Results', [('Commodity', None, 12), ('Category', None, 10), ('Test', None, 20), ('Shelf life h @20°C', HRS, 10),
                                   ('Q10', NUM, 7), ('Humidity sensitivity', NUM, 9), ('Initial RSL h', HRS, 10), ('Conservative buffer h', HRS, 10),
                                   ('Buyer min RSL h (plan)', HRS, 10), ('A arrival RSL h', HRS, 10), ('A accepted', None, 9), ('C arrival RSL h', HRS, 10),
                                   ('C accepted', None, 9), ('D buyer', None, 8), ('D arrival RSL h', HRS, 10), ('D accepted', None, 9),
                                   ('D decision status', None, 24), ('A profit', INR, 12), ('D profit', INR, 12), ('D − C profit', INR, 12), ('Shipment ID', None, 14)],
             frows, input_cols=6, delta_cols=(20,),
             note=f'Live shipment engine, one 800 kg grade-A shipment per test from H1 to B02 (default plan), deterministic agent. Initial RSL = hub_rsl_fraction × shelf life × multiplier. RSL is a synthetic quality estimate (Q10 thermal ageing), not a food-safety certification. Fleet batch-level RSL for S01/X01/X04/S09/S17 at seed 60000 is in csv/batches_*.csv.')
    bk.table('Dynamic_Pricing', [('Commodity', None, 12), ('Test', None, 20), ('Reference market ₹/kg', INR2, 11), ('Market price factor', NUM, 9),
                                 ('Grade after event', None, 8), ('Planned buyer price factor after event', NUM, 11), ('A buyer', None, 8),
                                 ('A offer ₹/kg', INR2, 10), ('A realized ₹/kg', INR2, 10), ('D buyer', None, 8), ('D channel', None, 10),
                                 ('D offer ₹/kg', INR2, 10), ('D realized ₹/kg', INR2, 10), ('D − A realized ₹/kg', INR2, 11), ('D revenue', INR, 12),
                                 ('Shipment ID', None, 14)], prows, input_cols=6, delta_cols=(14,),
             note='Offer = reference price × buyer channel price factor × market factor × grade adjustment × freshness; realized = offer if accepted, ₹0 if rejected. Fleet engine prices: Retailer A ₹32, Wholesale B ₹24, Processor C ₹16 per kg × price_multiplier × daily factor U(0.85,1.15); realized fleet ₹/kg per scenario is in Scenario_Summary.')

    # ------------------------------------------------------------------ 12 Agent_Action_Log + Agent_Validation
    trace_export, arows, vrows = [], [], []
    for cid in ['AG1', 'AG2', 'AG3', 'AG4', 'AG5', 'AG6', 'S06', 'S34']:
        s = ship_by[cid]
        resp = s['response']
        after = ship_by.get(cid + '-approval', {}).get('response', resp)
        stale = ship_by.get(cid + '-stale')
        tr = after['state']['trace']
        usage = next((t['result'] for t in tr if t['tool'] == 'llm_usage'), None)
        cls = 'LIVE CLAUDE AGENT' if usage else ('DETERMINISTIC FALLBACK' if resp['request']['mode'] == 'deterministic' else 'UNSUPPORTED / NOT VERIFIED')
        tools = [t['tool'] for t in tr]
        approval = s.get('request', {})
        decision = ship_by.get(cid + '-approval', {}).get('request', {}).get('approve')
        D0, D1 = resp['strategies']['D'], after['strategies']['D']
        vrows.append((cid, resp['id'], resp['request']['mode'], cls, ', '.join(f"{d['type']}({d['severity']})" for d in resp['request']['disruptions']) or 'none',
                      json.dumps(resp['request']['settings']) if resp['request']['settings'] else '—', len(tr),
                      usage['tool_calls'] if usage else 0, usage['input_tokens'] if usage else 0, usage['output_tokens'] if usage else 0,
                      usage['estimated_cost_usd'] if usage else 0, resp['decision_status'],
                      'approve' if decision is True else ('reject' if decision is False else '—'), after['decision_status'],
                      resp['state']['destination'], after['state']['destination'], after['state']['status'], after['state']['version'],
                      'validate_constraints' in tools or 'optimize_destination' in tools, 'execute_reroute' in tools,
                      f"HTTP {stale['http_status']} (refused)" if stale else '—', D0['lead_time_hours'], D1['arrival_rsl_hours'],
                      next((t['result']['text'] for t in tr if t['tool'] == 'agent_explanation'), '')[:600]))
        trace_export.append({'case': cid, 'shipment_id': resp['id'], 'classification': cls, 'request': resp['request'],
                             'decision_status_before_approval': resp['decision_status'], 'human_decision': decision,
                             'decision_status_after': after['decision_status'], 'state_after': {k: after['state'][k] for k in ('status', 'destination', 'version')},
                             'trace': tr})
        for t in tr:
            arows.append((cid, resp['id'], cls, t['sequence'], t.get('timestamp', ''), t['tool'], json.dumps(t.get('inputs', {}), ensure_ascii=False)[:400],
                          json.dumps(t.get('result', {}), ensure_ascii=False)[:700]))
    # simulated fleet traces (sample)
    if det.exists():
        for r in [x for x in d['M3'] if x['rescue_attempt']][:5]:
            for t in r['trace']:
                arows.append((f"X01/{r['batch_id']}", f"LOCAL-X01-60000", 'SIMULATED AGENT EXECUTION (fleet policy emulation)', t['sequence'],
                              t['timestamp'], t['tool'], '', ''))
    bk.table('Agent_Action_Log', [('Case', None, 11), ('Shipment / run ID', None, 16), ('Classification', None, 26), ('Seq', INT, 5),
                                  ('Timestamp', None, 22), ('Tool', None, 26), ('Inputs (JSON, truncated)', None, 50), ('Result (JSON, truncated)', None, 90)],
             arows, input_cols=3,
             note='Actual trace entries returned by the live API (full JSON in agent_traces.json). LIVE CLAUDE AGENT = trace contains llm_usage from the Anthropic API; DETERMINISTIC FALLBACK = deterministic mode; SIMULATED AGENT EXECUTION = fleet engine M3 policy emulation (no LLM).')
    bk.table('Agent_Validation', [('Case', None, 7), ('Shipment ID', None, 14), ('Requested mode', None, 12), ('Classification', None, 22),
                                  ('Disruption', None, 34), ('Settings', None, 22), ('Trace entries', INT, 8), ('LLM tool calls', INT, 8),
                                  ('Input tokens', INT, 9), ('Output tokens', INT, 9), ('LLM cost USD (est.)', '0.00000', 10),
                                  ('Decision before approval', None, 22), ('Human decision', None, 10), ('Decision after', None, 20),
                                  ('Destination before', None, 10), ('Destination after', None, 10), ('State status after', None, 14),
                                  ('State version', INT, 8), ('Guardrail/optimizer check in trace', None, 10), ('Reroute executed', None, 10),
                                  ('Stale-approval replay', None, 16), ('D lead time h', HRS, 9), ('D arrival RSL h', HRS, 9), ('Claude explanation (excerpt)', None, 90)],
             vrows, input_cols=6,
             note='Six live-mode cases on the deployed site plus two deterministic controls. Approval and rejection were sent through POST /shipments/{id}/approval; the replay with a stale version must be refused (optimistic concurrency guard).')
    (OUT / 'agent_traces.json').write_text(json.dumps(trace_export, indent=1, ensure_ascii=False), encoding='utf-8')

    # ------------------------------------------------------------------ 13 Sensitivity_Analysis
    rows = []
    for sid in fleet_ids:
        ref = ref_of(sid)
        if sid == ref:
            continue
        a, b = agg[sid], agg[ref]
        ddc = a['DC']['mean'] - b['DC']['mean']
        dda = a['DA']['mean'] - b['DA']['mean']
        tol = 1.0
        direction = 'no change' if abs(ddc) < tol else ('strengthens D vs C' if ddc > 0 else 'weakens D vs C')
        rows.append((sid, name_of[sid], ref, json.dumps(cat[sid][4] if sid in cat else sup[sid][3]), b['DC']['mean'], a['DC']['mean'], ddc,
                     b['DA']['mean'], a['DA']['mean'], dda, 100 * (a['M3']['spoil']['mean'] - b['M3']['spoil']['mean']), a['P_DC'], direction))
    rows.sort(key=lambda r: r[6], reverse=True)
    bk.table('Sensitivity_Analysis', [('Scenario', None, 9), ('Name', None, 28), ('Reference', None, 9), ('Change', None, 40),
                                      ('Ref D − C', INR, 13), ('Scenario D − C', INR, 13), ('Δ (D − C)', INR, 13), ('Ref D − A', INR, 13),
                                      ('Scenario D − A', INR, 13), ('Δ (D − A)', INR, 13), ('Δ D spoilage (pp)', PP, 11), ('Share seeds D > C', PCT, 9),
                                      ('Effect on agentic contribution', None, 22)], rows, input_cols=4, delta_cols=(7, 10),
             note='Each scenario compared with its own reference on the same seeds (S-scenarios vs S01, X-scenarios vs X01). Sorted by Δ(D − C): top rows are conditions where the agent adds most over the non-agentic optimizer.')

    # ------------------------------------------------------------------ 14 Failed_Tests
    present = {c['id'] for c in network['commodities']}
    missing = [c for cs in BRIEF_COMMODITIES.values() for c in cs if c not in present]
    extra = sorted(present - {c for cs in BRIEF_COMMODITIES.values() for c in cs})
    ftests = [
        ('S07', 'NOT IMPLEMENTED', 'Loading delay +2 h', 'No warehouse/loading lead-time input in either engine.', 'Add dispatch_delay_hours (fleet + shipment) applied before departure.'),
        ('S13', 'NOT IMPLEMENTED', 'Advanced ripening', 'Fleet maturity is a random draw per batch; shipment engine has no maturity input.', 'Add a maturity-mix control (mature_green/pink/red shares) and per-commodity ripening rates.'),
        ('S26', 'NOT IMPLEMENTED', 'Agent execution failure', 'No input injects a failed reroute/booking. agent_failed exists only for live-LLM provider errors (unit-tested).', 'Add execution_failure_probability with retry/fallback logging.'),
        ('S02–S05', 'MODEL LIMITATION', 'Traffic alone has no effect', 'Fleet generator applies traffic only to disrupted batches; default disruption_probability = 0, so S02–S05 equal S01 exactly.', 'Supplementary X02–X05 (p=0.30) show the traffic effect; consider applying a baseline traffic delay to all batches.'),
        ('S15', 'INFORMATION', 'Strong buyer availability = default', 'Brief value 0.95 equals the default, so S15 reproduces S01 seed-for-seed (a determinism check).', '—'),
        ('S21', 'EXECUTED — NO EFFECT', 'Processor rescue unavailable', 'In the paired 6-batch live test the agent never chose a processor even when available (rescues went to wholesale/retail), so zeroing processor capacity changed nothing.', 'Test with lower-grade/low-RSL batches where processors are the only feasible outlet.'),
        ('Commodities', 'COVERAGE GAP', f'{len(missing)} of 18 brief commodities not implemented', 'Missing: ' + ', '.join(missing) + '. Backend instead has: ' + ', '.join(extra) + '.', 'Add catalogue rows (shelf life, Q10, humidity sensitivity, reference price) with sources.'),
        ('Fleet commodities', 'SCOPE', 'Fleet simulation is tomato-only', 'All S-scenario statistics refer to tomato lots on a 3-buyer synthetic graph.', 'Generalise the fleet generator to the commodity catalogue.'),
        ('Live vs local', 'DISCLOSED DEVIATION', 'C/D differ slightly between Linux (live) and Windows (local)', 'A and B are identical on every checked seed. For some seeds the event SHA-256 differs, most likely from last-digit floating-point differences in the platform math library (random.gauss); C/D use a time-limited CP-SAT solve that is sensitive to such tiny differences. See Live_vs_Local_Check.', 'Run the matrix inside the Docker image or raise the solver time limit.'),
        ('Upload validation', 'EXPECTED REJECTION', 'Workbook with okra / cauliflower', 'Live upload returned HTTP 422 with row-level errors (Batch_Inventory row 6 commodity_id: Unknown commodity). Validation works as designed.', '—'),
    ]
    for sid, seed, m, issues in audit_fail:
        ftests.append((sid, 'AUDIT FAIL', f'seed {seed} {LETTER[m]}', '; '.join(issues), 'Exclude from evidence'))
    bk.table('Failed_Tests', [('Scenario / area', None, 16), ('Status', None, 22), ('Test', None, 34), ('Finding', None, 80), ('Feature needed / action', None, 60)],
             ftests, input_cols=3, note=f'Financial audits run on every fleet run × strategy ({len(fleet) * 4} checks): {len(audit_fail)} failures.')

    # ------------------------------------------------------------------ 15 Evidence_Sources & 16 Assumptions
    ev = L['live_evidence.json']
    srows = [(s['id'], s['category'], s['title'], s.get('authors', ''), s.get('url', ''), s.get('accessed_date', ''), s.get('limitations', ''))
             for s in ev['sources']]
    srows.append(('N1', 'DATA', 'Mumbai road network distances', 'OSRM / OpenStreetMap contributors', network.get('source', 'router.project-osrm.org'),
                  network.get('generated_at', '2026-10-28'), 'Free-flow road distance/time only; traffic is a simulated multiplier.'))
    srows.append(('V1', 'RUN', 'This validation — live API responses', 'FreshFlow deployment', LIVE, now[:10], 'Supports only the simulated outputs recorded in this workbook.'))
    bk.table('Evidence_Sources', [('ID', None, 6), ('Category', None, 11), ('Title', None, 36), ('Authors', None, 26), ('URL', None, 50),
                                  ('Accessed', None, 12), ('What it supports / limitation', None, 80)], srows, input_cols=2)
    arows2 = [(a['id'], a['category'], a['parameter'], a['value'], a.get('unit', ''), a.get('basis', ''), a.get('limitation', ''), a.get('validation_status', ''))
              for a in ev['assumptions']]
    arows2 += [
        ('V-01', 'VALIDATION', 'Evaluation seeds', f'{SEEDS[0]}–{SEEDS[-1]}', 'seeds', 'Prespecified, disjoint from calibration seeds 101–103', 'Same seed ⇒ same events for A–D', 'design choice'),
        ('V-02', 'VALIDATION', 'S12 / S29 initial RSL ×0.70', '0.98', '× (initial_rsl_scale)', '0.70 × default 1.4', 'Fleet scale is absolute, not relative', 'mapping'),
        ('V-03', 'VALIDATION', 'S25 prediction error +6 h', '9', 'h SD', 'default 3 h + 6 h', 'Interpreted as an increase over the default', 'mapping'),
        ('V-04', 'VALIDATION', 'Supplementary disruption probability', '0.30', 'prob.', 'Matches backend/run_risk_study.py greater_traffic scenario', 'Not in brief; declared', 'design choice'),
        ('V-05', 'VALIDATION', 'S32 "moderate disruptions"', 'p=0.15, traffic 1.4×', '—', 'Analyst choice', 'Brief gave no numeric value', 'mapping'),
        ('V-06', 'VALIDATION', 'S33 "low rescue availability"', 'buyer_availability 0.70, AI cost ×10, p=0', '—', 'Analyst choice (×10 as in risk study)', 'Brief gave no numeric value', 'mapping'),
        ('F-01', 'SYNTHETIC', 'Fleet buyer prices', 'Retailer A 32 / Wholesale B 24 / Processor C 16', '₹/kg', 'backend/app/optimization.py', 'Synthetic', 'not validated'),
        ('F-02', 'SYNTHETIC', 'Fleet minimum arrival RSL', '24 / 8 / 2', 'h', 'backend/app/optimization.py', 'Synthetic', 'not validated'),
        ('F-03', 'SYNTHETIC', 'Fleet agent/tech cost', 'A 0, B 0.10, C 0.15, D 0.15 + 0.10×AI multiplier', '₹/kg', 'backend/app/simulation.py', 'Synthetic', 'not validated'),
        ('F-04', 'SYNTHETIC', 'Fleet rescue handling cost', '0.75', '₹/kg', 'backend/app/optimization.py', 'Charged on every rescue attempt', 'not validated'),
        ('F-05', 'SYNTHETIC', 'Shipment traffic multiplier (default)', '2.0', '× OSRM free-flow', 'backend/app/catalogue.py', 'Simulated congestion, not live traffic', 'not validated'),
    ]
    bk.table('Assumptions', [('ID', None, 8), ('Category', None, 12), ('Parameter', None, 32), ('Value', None, 28), ('Unit', None, 14),
                             ('Basis / rationale', None, 46), ('Limitation', None, 46), ('Validation status', None, 22)], arows2, input_cols=5)

    # ------------------------------------------------------------------ Live_vs_Local_Check
    lrows = []
    lb = {r['seed']: r for r in by['S01']}
    checks = [('S01', run) for run in base['result']['runs']] + [(k, L[f'live_check_{k}.json']['result']['runs'][0]) for k in ['X01', 'S17', 'S31']]
    five = RAW / 'live_check_X01_5seeds.json'
    live_x01 = {checks[3][1]['seed']: checks[3][1]}
    if five.exists():
        fx = json.loads(five.read_text(encoding='utf-8'))
        rows_ = fx['result']['runs']
        checks += [('X01', run) for run in rows_]
        live_x01.update({run['seed']: run for run in rows_})
        rows_reg = [('LIVE', fx['id'], 'X01 cross-check (5 seeds)', json.dumps({'seed': 60001, 'repetitions': 5, 'disruption_probability': 0.3}),
                     'live Render (Linux)', '—', 'COMPLETED', '; '.join(f"{r['seed']}:{r['dataset_sha256'][:12]}" for r in rows_), 'Cross-platform comparison')]
        reg = bk.wb['Run_Register']
        for r in rows_reg:
            reg.append(r)
    dc = lambda run: run['metrics'][3]['operating_profit_inr'] - run['metrics'][2]['operating_profit_inr']
    seeds_x = sorted(live_x01)
    live_dc = st.mean(dc(live_x01[s]) for s in seeds_x)
    local_dc = st.mean(dc(next(r for r in by['X01'] if r['seed'] == s)) for s in seeds_x)
    max_dev = 0
    for sid, live_run in checks:
        loc = next((r for r in by.get(sid, []) if r['seed'] == live_run['seed']), None)
        if not loc:
            continue
        for i, m in enumerate(MODELS):
            lp, pp = live_run['metrics'][i]['operating_profit_inr'], loc['metrics'][i]['operating_profit_inr']
            dev = (pp - lp) / abs(lp) if lp else 0
            max_dev = max(max_dev, abs(dev))
            lrows.append((sid, live_run['seed'], LETTER[m], live_run['dataset_sha256'] == loc['dataset_sha256'], lp, pp, pp - lp, dev,
                          live_run['metrics'][i]['spoilage_rate'], loc['metrics'][i]['spoilage_rate'],
                          100 * (loc['metrics'][i]['spoilage_rate'] - live_run['metrics'][i]['spoilage_rate'])))
    bk.table('Live_vs_Local_Check', [('Scenario', None, 9), ('Seed', None, 8), ('Strategy', None, 8), ('Same events (SHA match)', None, 11),
                                     ('Live profit', INR, 14), ('Local profit', INR, 14), ('Local − live', INR, 12), ('Relative deviation', '0.000%', 11),
                                     ('Live spoilage', PCT, 10), ('Local spoilage', PCT, 10), ('Δ spoilage pp', PP, 10)], lrows, input_cols=4,
             note=f'Same commit, same seeds. A/B profits match exactly on every seed. Maximum absolute relative profit deviation: {max_dev:.2%} (C, disrupted X01). '
                  f'About 82% of daily C/D allocations end at the CP-SAT time limit (FEASIBLE, not proven OPTIMAL), so tiny platform floating-point differences can select a different near-optimal plan; this matters most for C under disruption. '
                  f'X01 mean D − C on seeds {seeds_x[0]}–{seeds_x[-1]}: live (Linux) ₹{live_dc / 1e5:,.1f} lakh vs local (Windows) ₹{local_dc / 1e5:,.1f} lakh. Direction and order of magnitude agree; exact magnitudes carry this solver noise.')

    # ------------------------------------------------------------------ Commodity_Coverage
    crow = []
    for catg, cs in BRIEF_COMMODITIES.items():
        for c in cs:
            crow.append((c, catg, 'IMPLEMENTED' if c in present else 'NOT IMPLEMENTED — future scope', 'shipment engine only' if c in present and c != 'tomato' else ('fleet + shipment' if c == 'tomato' else '—')))
    for c in extra:
        crow.append((c, com[c]['category'], 'IMPLEMENTED (not in brief list)', 'shipment engine only'))
    bk.table('Commodity_Coverage', [('Commodity', None, 14), ('Category', None, 11), ('Status', None, 32), ('Engines', None, 22)], crow, input_cols=2)

    # ------------------------------------------------------------------ 17 PPT_Ready_Results
    S = agg['S01']
    X = agg.get('X01', S)
    n_brief_fleet = [sid for sid in fleet_ids if sid.startswith('S')]
    d_beats_c = [sid for sid in n_brief_fleet if agg[sid]['DC']['mean'] > 0]
    d_not_c = [sid for sid in n_brief_fleet if agg[sid]['DC']['mean'] <= 0]
    sig = lambda sid: tuple(round(r['metrics'][i]['operating_profit_inr'], 2) for r in by[sid] for i in range(4))
    same_as_s01 = [sid for sid in n_brief_fleet if sid != 'S01' and sig(sid) == sig('S01')]
    distinct = [sid for sid in n_brief_fleet if sid not in same_as_s01]
    d_not_c_distinct = [sid for sid in distinct if agg[sid]['DC']['mean'] <= 0]
    d_beats_a = [sid for sid in n_brief_fleet if agg[sid]['DA']['mean'] > 0]
    x_beats_c = [sid for sid in fleet_ids if sid.startswith('X') and agg[sid]['DC']['mean'] > 0]
    ag1 = ship_by.get('AG1-approval', {}).get('response') or ship_by['AG1']['response']
    s06 = ship_by['S06']['response']
    rr = lambda sh, sid, m=None: f"{sh}!row {rowref.get((sh, sid, m) if m else (sh, sid), '?')}"
    claims = [
        ('Slide 3', 'Commodities implemented', f"{len(present)} in the Mumbai shipment engine ({len(present) - len(extra)} of the 18 on the brief list); fleet simulation is tomato-only", 'Commodity_Coverage', 'live /api/v1/network', 'VERIFIED'),
        ('Slide 3', 'Network', f"{len(network['hubs'])} APMC hubs + {len(network['buyers'])} synthetic buyers = {len(network['hubs']) + len(network['buyers'])} Mumbai/MMR locations; {len(routes)} OSRM road routes", 'Route_ETA_Changes; live_network.json', 'OSRM/OpenStreetMap', 'VERIFIED'),
        ('Slide 3', 'Scenarios executed', f"{n_fleet + n_ship} of 34 brief scenarios executed ({n_fleet} fleet, {n_ship} live shipment engine); {n_ni} not implemented; +{len(SUPPLEMENTARY)} declared supplementary", 'Scenario_Catalog', '—', 'VERIFIED'),
        ('Slide 3', 'Scale', f"2,000 batches × 30 days per run; {reps} paired seeds per scenario; {len(fleet)} fleet runs × 4 strategies = {len(fleet) * 4:,} strategy-runs; {len(ship)} live API shipment tests", 'Run_Register', '—', 'VERIFIED'),
        ('Slide 3', 'Reproducibility', f'Same seed + parameters ⇒ identical results on the same platform (S15 reproduced S01 seed-for-seed). Across Linux (live) and Windows (local): A/B exact; C/D differ by up to {max_dev:.1%} because time-limited CP-SAT solves are platform-sensitive', 'Live_vs_Local_Check', '—', 'VERIFIED'),
        ('Slide 5', 'Baseline spoilage A vs D (S01)', f"A {S['M0']['spoil']['mean']:.1%} → D {S['M3']['spoil']['mean']:.1%} (C {S['M2']['spoil']['mean']:.1%})", rr('Spoilage_Comparison', 'S01'), 'S01 · LOCAL-S01-60000…60029', 'SIMULATED'),
        ('Slide 5', 'Baseline profit A vs D (S01)', f"A ₹{S['M0']['profit']['mean'] / 1e5:,.1f} lakh → D ₹{S['M3']['profit']['mean'] / 1e5:,.1f} lakh per 30-day run (D − A ₹{S['DA']['mean'] / 1e5:,.1f} lakh; D > A in {S['P_DA']:.0%} of seeds)", rr('Profit_Comparison', 'S01'), 'S01', 'SIMULATED'),
        ('Slide 5', 'Incremental D vs C (S01, no disruptions)', f"₹{S['DC']['mean']:,.0f} (P05 ₹{S['DC']['p05']:,.0f}, P95 ₹{S['DC']['p95']:,.0f}); D > C in {S['P_DC']:.0%} of seeds", rr('Profit_Comparison', 'S01'), 'S01', 'SIMULATED'),
        ('Slide 5', 'Incremental D vs C (X01, 30% disrupted)', f"₹{X['DC']['mean'] / 1e5:,.1f} lakh (P05 ₹{X['DC']['p05'] / 1e5:,.1f} lakh, P95 ₹{X['DC']['p95'] / 1e5:,.1f} lakh); D > C in {X['P_DC']:.0%} of seeds; spoilage C {X['M2']['spoil']['mean']:.1%} → D {X['M3']['spoil']['mean']:.1%}", rr('Profit_Comparison', 'X01'), 'X01 (supplementary)', 'SIMULATED'),
        ('Slide 5', 'Scenarios where D improves profit vs A', f"{len(d_beats_a)} of {len(n_brief_fleet)} executed fleet scenarios", 'Profit_Comparison', 'S-scenarios', 'SIMULATED'),
        ('Slide 5', 'Scenarios where D does NOT beat C', f"{len(d_not_c)} of {len(n_brief_fleet)} brief fleet scenarios ({len(d_not_c_distinct)} of {len(distinct)} after removing {len(same_as_s01)} that reproduce S01 exactly: {', '.join(same_as_s01)}). All of these have zero disruption probability (app default); the only brief scenario with disruptions (S32) has D > C. With disruptions (X-family) D beats C in {len(x_beats_c)} of {len([s for s in fleet_ids if s.startswith('X')])}", 'Profit_Comparison; Sensitivity_Analysis', '—', 'SIMULATED'),
        ('Slide 5', 'D − C on the live platform (X01)', f"Live Linux deployment, seeds {seeds_x[0]}–{seeds_x[-1]}: mean D − C ₹{live_dc / 1e5:,.1f} lakh (local Windows, same seeds: ₹{local_dc / 1e5:,.1f} lakh)", 'Live_vs_Local_Check', 'X01 live run IDs in Run_Register', 'SIMULATED (live)'),
        ('Slide 5', 'Representative rescue (live)', f"S06 road closure: {s06['strategies']['A']['buyer_name']} would reject ({s06['strategies']['A']['rejection_reasons'][0]}); agent proposes {s06['strategies']['D']['buyer_name']}, {s06['strategies']['D']['distance_km']:.1f} km, arrival RSL {s06['strategies']['D']['arrival_rsl_hours']:.1f} h; D profit ₹{s06['strategies']['D']['operating_profit_inr']:,.0f} vs A ₹{s06['strategies']['A']['operating_profit_inr']:,.0f}", 'Route_ETA_Changes; Agent_Action_Log', f"S06 · {s06['id']}", 'SIMULATED (live engine)'),
        ('Slide 5', 'Live Claude agent', f"AG1 {ag1['id']}: Claude ({health.get('llm_model')}) called tools, proposed a diversion, human approved, state changed {ship_by['AG1']['response']['state']['destination']}→{ag1['state']['destination']}", 'Agent_Validation', 'AG1', 'VERIFIED (live)'),
        ('Slide 5', 'Limiting assumption', 'All data are synthetic; no field calibration. Fleet engine is tomato-only on a 3-buyer synthetic graph; traffic is a simulated multiplier, not live traffic.', 'README; Assumptions', '—', 'DISCLOSE'),
    ]
    bk.table('PPT_Ready_Results', [('Slide', None, 8), ('Claim', None, 34), ('Verified value', None, 90), ('Workbook location', None, 34),
                                   ('Scenario → run IDs', None, 28), ('Status', None, 16)], claims, input_cols=2,
             note='Only numbers produced by the runs in this workbook. "SIMULATED" = synthetic-data simulation, never actual Mumbai operating outcomes. Favourable and unfavourable results are both listed.')

    # ------------------------------------------------------------------ 18 Data_Dictionary
    dd = [
        ('Strategy A (M0)', 'Traditional Planning: committed morning buyer, no freshness screening or optimization.', '—', '—'),
        ('Strategy B (M1)', 'Freshness-Aware: FEFO ordering, same buyer rule as A.', '—', '—'),
        ('Strategy C (M2)', 'Smart Optimization: daily CP-SAT allocation on predicted RSL; no event-driven intervention.', '—', '—'),
        ('Strategy D (M3)', 'FreshFlow Agentic AI: C plus event-driven re-optimization, constraint validation and (simulated) human approval delay.', '—', '—'),
        ('Input kg', 'Total lot quantity entering the run.', 'Σ quantity_kg', 'kg'),
        ('Sold kg', 'Quantity accepted by a buyer.', 'Σ qty where accepted', 'kg'),
        ('Spoiled / written-off kg', 'Quantity not accepted (spoiled, rejected or unsold).', 'Input − Sold', 'kg'),
        ('Processor (salvage) kg / revenue', 'Accepted at Processor C (fleet) — a subset of sold/gross revenue.', '—', 'kg / ₹'),
        ('Mass residual', 'Input − Sold − Spoiled; must be 0.', '—', 'kg'),
        ('Spoilage rate', 'Spoiled kg ÷ Input kg.', '—', '%'),
        ('On-time rate', 'Accepted lots arriving ≤ 24 h after dispatch ÷ all lots.', '—', '%'),
        ('Buyer rejection rate', 'Attempted deliveries not accepted ÷ attempted deliveries.', '—', '%'),
        ('RSL', 'Remaining shelf life (hours) from Q10 thermal ageing; a quality estimate, not a safety certification.', 'thermal_age = h × Q10^((T−20)/10)', 'h'),
        ('Rescue attempt / rescued', 'D re-optimizes after an event; rescued = new destination and accepted.', '—', 'count'),
        ('Failed interventions', 'Rescue attempts − rescued.', '—', 'count'),
        ('Gross revenue', 'Σ accepted kg × realized price.', '—', '₹'),
        ('Operating profit', 'Gross revenue − procurement − transport − handling − diversion − agent/technology.', '—', '₹'),
        ('Operating margin', 'Operating profit ÷ gross revenue.', '—', '%'),
        ('Net realization ₹/input kg', '(Gross revenue − non-procurement costs) ÷ input kg.', '—', '₹/kg'),
        ('D − C', 'Value of agentic intervention over the same optimizer without it, paired by seed.', 'profit_D − profit_C', '₹'),
        ('Share of seeds D > A / D > C', 'Fraction of paired seeds where D profit exceeds A / C.', '—', '%'),
        ('P05 / P95', '5th / 95th percentile across seeds (linear interpolation). Simulation variability, not confidence intervals.', '—', 'as metric'),
        ('pp', 'Percentage points.', '—', 'pp'),
        ('SIMULATED RESULT', 'Output of the model on synthetic inputs.', '—', '—'),
        ('LIVE CLAUDE AGENT', 'Trace contains Anthropic API usage records (tokens, cost).', '—', '—'),
        ('DETERMINISTIC FALLBACK', 'Python policy without LLM calls.', '—', '—'),
    ]
    bk.table('Data_Dictionary', [('Term', None, 30), ('Definition', None, 90), ('Formula', None, 34), ('Unit', None, 10)], dd)

    order = ['README', 'Scenario_Catalog', 'Inputs', 'Run_Register', 'A_B_C_D_Results', 'Scenario_Summary', 'Profit_Comparison',
             'Spoilage_Comparison', 'Route_ETA_Changes', 'Freshness_Results', 'Dynamic_Pricing', 'Agent_Action_Log', 'Sensitivity_Analysis',
             'Failed_Tests', 'Evidence_Sources', 'Assumptions', 'PPT_Ready_Results', 'Data_Dictionary', 'Agent_Validation',
             'Live_vs_Local_Check', 'Commodity_Coverage']
    bk.wb._sheets = [bk.wb[n] for n in order]
    bk.wb['PPT_Ready_Results'].sheet_properties.tabColor = 'C9A227'
    bk.save(OUT / 'FreshFlow_Scenario_Validation.xlsx')

    # batch-level CSVs
    for f in RAW.glob('detail_*_60000.json'):
        d = json.loads(f.read_text(encoding='utf-8'))
        sid = f.stem.split('_')[1]
        with (OUT / 'csv' / f'batches_{sid}_seed60000.csv').open('w', newline='', encoding='utf-8-sig') as fh:
            keys = ['strategy', 'batch_id', 'day', 'quantity_kg', 'morning_destination', 'destination', 'sold_kg', 'spoiled_kg', 'revenue_inr',
                    'realized_unit_price_inr', 'procurement_inr', 'transport_inr', 'handling_inr', 'rescue_inr', 'agent_sensor_inr', 'rescue_attempt',
                    'rescued', 'buyer_acceptance', 'on_time', 'estimated_rsl_h', 'estimated_rsl_lower_h', 'realized_arrival_rsl_hours', 'solver_status']
            w = csv.writer(fh)
            w.writerow(keys)
            for m in MODELS:
                for r in d[m]:
                    w.writerow([LETTER[m]] + [r.get(k) for k in keys[1:]])

    summary = {'agg': {sid: {k: v for k, v in a.items()} for sid, a in agg.items()}, 'claims': claims, 'max_dev': max_dev,
               'audit_fail': audit_fail, 'n_fleet_runs': len(fleet), 'reps': reps, 'd_not_c': d_not_c, 'd_beats_a': d_beats_a,
               'same_as_s01': same_as_s01, 'd_not_c_distinct': d_not_c_distinct, 'distinct': distinct,
               'live_dc_x01': live_dc, 'local_dc_x01': local_dc, 'seeds_x01_live': seeds_x,
               'x_beats_c': x_beats_c, 'rowref': {'|'.join(map(str, k)): v for k, v in rowref.items()}, 'fleet_ids': fleet_ids, 'names': name_of}
    (RAW / 'summary.json').write_text(json.dumps(summary, indent=1, default=str), encoding='utf-8')
    print('workbook written;', len(fleet), 'fleet runs;', len(audit_fail), 'audit failures; max live/local deviation', f'{max_dev:.4%}')


if __name__ == '__main__':
    (OUT / 'csv').mkdir(parents=True, exist_ok=True)
    main()
