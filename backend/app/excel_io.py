"""Excel input template, validated upload and result exports (openpyxl).

Input workbook  -> dataset used by the shipment engine (every Batch_Inventory row is run through A–D).
Result exports  -> actual backend outputs with run ids and seeds, so they can be reproduced.
"""
import base64
import copy
import csv
import io
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .catalogue import DISRUPTIONS, SETTING_BOUNDS, SEVERITIES

RESEARCH = Path(__file__).resolve().parents[2] / 'research'
HEADER_FILL = PatternFill('solid', fgColor='1D4433')
INPUT_FILL = PatternFill('solid', fgColor='EEF5F1')
HEADER_FONT = Font(bold=True, color='FFFFFF')

# sheet -> (columns, description). Units are in column names.
INPUT_SHEETS = {
    'Locations': (['location_id', 'name', 'type', 'channel', 'area', 'lat', 'lon'], 'Hubs (type=hub) and buyers (type=buyer).'),
    'Road_Routes': (['origin_id', 'buyer_id', 'distance_km', 'freeflow_minutes', 'source'], 'Hub → buyer road routes (OSRM over OpenStreetMap). Diversion legs from checkpoints use the built-in network.'),
    'Commodities': (['commodity_id', 'name', 'category', 'reference_price_inr_kg'], 'Produce and illustrative reference price.'),
    'Batch_Inventory': (['batch_id', 'commodity_id', 'origin_id', 'committed_buyer_id', 'quantity_kg', 'quality_grade'], 'One row per shipment to simulate.'),
    'Buyer_Offers': (['buyer_id', 'price_factor_vs_reference', 'min_rsl_fraction_of_shelf_life', 'accepted_grades'], 'Buyer offer = reference price × factor × grade × freshness factor.'),
    'Buyer_Capacity': (['buyer_id', 'capacity_kg'], 'Maximum kg a buyer accepts for one shipment.'),
    'Freshness_Parameters': (['commodity_id', 'shelf_life_hours_20c', 'q10', 'humidity_sensitivity'], 'Thermal-age proxy parameters (synthetic).'),
    'Disruption_Events': (['batch_id', 'disruption_type', 'severity'], 'Allowed types: ' + ', '.join(DISRUPTIONS) + '. Severity: low, medium, high.'),
    'Operating_Costs': (['cost_item', 'value', 'unit'], 'Logistics, handling, diversion and technology costs.'),
    'Scenario_Settings': (['setting', 'value', 'min', 'max'], 'Engine settings within the shown bounds.'),
    'Sources_Assumptions': (['id', 'category', 'item', 'value', 'basis_or_url', 'limitation'], 'Reference only — not read on upload.'),
}
COST_UNITS = {'truck_fixed_inr': 'INR/trip', 'truck_per_km_inr': 'INR/km', 'handling_inr_kg': 'INR/kg', 'diversion_fixed_inr': 'INR/diversion',
              'procurement_share_of_reference': 'fraction', 'tech_inr_kg_A': 'INR/kg', 'tech_inr_kg_B': 'INR/kg', 'tech_inr_kg_C': 'INR/kg',
              'tech_inr_kg_D_base': 'INR/kg', 'tech_inr_kg_D_agent': 'INR/kg × agent cost multiplier'}


def style(ws, header_row=1, input_sheet=False):
    for cell in ws[header_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for i, column in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    if input_sheet:
        ws.sheet_properties.tabColor = '0F8A7E'


def add_sheet(wb, name, columns, rows, input_sheet=False):
    ws = wb.create_sheet(name[:31])
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c) if isinstance(r, dict) else v for c, v in zip(columns, r if not isinstance(r, dict) else columns)])
    style(ws, input_sheet=input_sheet)
    return ws


def to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sources_rows():
    rows = []
    for name in ['assumptions', 'sources']:
        path = RESEARCH / f'{name}.csv'
        if path.exists():
            for r in csv.DictReader(path.open(encoding='utf-8')):
                rows.append({'id': r.get('id'), 'category': r.get('category'), 'item': r.get('parameter') or r.get('title'),
                             'value': r.get('value', ''), 'basis_or_url': r.get('basis') or r.get('url'),
                             'limitation': r.get('limitation') or r.get('limitations')})
    rows.append({'id': 'N1', 'category': 'SYNTHETIC', 'item': 'Mumbai network', 'value': '3 hubs + 12 synthetic buyers',
                 'basis_or_url': 'Approximate public coordinates; road routes from OSRM/OpenStreetMap', 'limitation': 'Traffic is simulated, not live'})
    return rows


# ----------------------------------------------------------------------------------- template

def template(ds, example_batches=None):
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet('README')
    readme.append(['FreshFlow input workbook'])
    readme.append([f'Generated {datetime.now(timezone.utc).isoformat(timespec="seconds")} from dataset "{ds["id"]}". All values are SYNTHETIC.'])
    readme.append(['Green tabs are inputs. Edit values, keep the column headers, upload in Data & Reports.'])
    for name, (_, desc) in INPUT_SHEETS.items():
        readme.append([name, desc])
    readme.column_dimensions['A'].width = 24
    readme.column_dimensions['B'].width = 120
    readme['A1'].font = Font(bold=True, size=14, color='1D4433')

    locs = [{'location_id': k, **{c: v.get(c) for c in ['name', 'type', 'channel', 'area', 'lat', 'lon']}} for k, v in ds['locations'].items()]
    add_sheet(wb, 'Locations', INPUT_SHEETS['Locations'][0], locs, True)
    routes = [{'origin_id': k.split('>')[0], 'buyer_id': k.split('>')[1], **v, 'source': 'OSRM/OpenStreetMap'} for k, v in ds['routes'].items()]
    add_sheet(wb, 'Road_Routes', INPUT_SHEETS['Road_Routes'][0], routes, True)
    add_sheet(wb, 'Commodities', INPUT_SHEETS['Commodities'][0],
              [{'commodity_id': c['id'], **c} for c in ds['commodities'].values()], True)
    batches = example_batches or [
        {'batch_id': 'B-001', 'commodity_id': 'tomato', 'origin_id': 'H1', 'committed_buyer_id': 'B02', 'quantity_kg': 800, 'quality_grade': 'A'},
        {'batch_id': 'B-002', 'commodity_id': 'spinach', 'origin_id': 'H1', 'committed_buyer_id': 'B05', 'quantity_kg': 400, 'quality_grade': 'A'},
        {'batch_id': 'B-003', 'commodity_id': 'mango', 'origin_id': 'H2', 'committed_buyer_id': 'B09', 'quantity_kg': 500, 'quality_grade': 'A'},
    ]
    add_sheet(wb, 'Batch_Inventory', INPUT_SHEETS['Batch_Inventory'][0], batches, True)
    add_sheet(wb, 'Buyer_Offers', INPUT_SHEETS['Buyer_Offers'][0],
              [{'buyer_id': k, 'price_factor_vs_reference': v['price_factor'], 'min_rsl_fraction_of_shelf_life': v['min_rsl_fraction'],
                'accepted_grades': v['accepted_grades']} for k, v in ds['buyers'].items()], True)
    add_sheet(wb, 'Buyer_Capacity', INPUT_SHEETS['Buyer_Capacity'][0],
              [{'buyer_id': k, 'capacity_kg': v['capacity_kg']} for k, v in ds['buyers'].items()], True)
    add_sheet(wb, 'Freshness_Parameters', INPUT_SHEETS['Freshness_Parameters'][0],
              [{'commodity_id': c['id'], **c} for c in ds['commodities'].values()], True)
    add_sheet(wb, 'Disruption_Events', INPUT_SHEETS['Disruption_Events'][0], [
        {'batch_id': 'B-001', 'disruption_type': 'heavy_traffic', 'severity': 'high'},
        {'batch_id': 'B-001', 'disruption_type': 'temperature_rise', 'severity': 'high'},
        {'batch_id': 'B-002', 'disruption_type': 'road_closure', 'severity': 'medium'},
    ], True)
    add_sheet(wb, 'Operating_Costs', INPUT_SHEETS['Operating_Costs'][0],
              [{'cost_item': k, 'value': v, 'unit': COST_UNITS.get(k, '')} for k, v in ds['costs'].items()], True)
    add_sheet(wb, 'Scenario_Settings', INPUT_SHEETS['Scenario_Settings'][0],
              [{'setting': k, 'value': v, 'min': SETTING_BOUNDS[k][0], 'max': SETTING_BOUNDS[k][1]} for k, v in ds['settings'].items()], True)
    add_sheet(wb, 'Sources_Assumptions', INPUT_SHEETS['Sources_Assumptions'][0], sources_rows())
    return to_bytes(wb)


# ----------------------------------------------------------------------------------- upload

class WorkbookErrors(Exception):
    def __init__(self, errors):
        super().__init__(f'{len(errors)} validation error(s)')
        self.errors = errors


def _rows(wb, sheet, errors):
    if sheet not in wb.sheetnames:
        errors.append({'sheet': sheet, 'row': None, 'column': None, 'message': 'Required sheet is missing'})
        return []
    ws = wb[sheet]
    values = list(ws.iter_rows(values_only=True))
    if not values:
        errors.append({'sheet': sheet, 'row': 1, 'column': None, 'message': 'Sheet is empty'})
        return []
    header = [str(h).strip() if h is not None else '' for h in values[0]]
    required = INPUT_SHEETS[sheet][0]
    missing = [c for c in required if c not in header]
    if missing:
        errors.append({'sheet': sheet, 'row': 1, 'column': ', '.join(missing), 'message': 'Missing required column(s)'})
        return []
    out = []
    for i, raw in enumerate(values[1:], start=2):
        if all(v is None or str(v).strip() == '' for v in raw):
            continue
        out.append((i, {h: raw[j] if j < len(raw) else None for j, h in enumerate(header)}))
    return out


def _num(errors, sheet, row, col, value, lo=None, hi=None, positive=False):
    try:
        x = float(value)
    except (TypeError, ValueError):
        errors.append({'sheet': sheet, 'row': row, 'column': col, 'message': f'"{value}" is not a number'})
        return None
    if x != x or (positive and x <= 0) or (lo is not None and x < lo) or (hi is not None and x > hi):
        bound = 'greater than 0' if positive else f'between {lo} and {hi}' if hi is not None else f'at least {lo}'
        errors.append({'sheet': sheet, 'row': row, 'column': col, 'message': f'{x} must be {bound}'})
        return None
    return x


def _unique(errors, sheet, rows, key):
    seen = {}
    for row, r in rows:
        k = str(r.get(key) or '').strip()
        if not k:
            errors.append({'sheet': sheet, 'row': row, 'column': key, 'message': 'Value is required'})
        elif k in seen:
            errors.append({'sheet': sheet, 'row': row, 'column': key, 'message': f'Duplicate id "{k}" (also row {seen[k]})'})
        else:
            seen[k] = row
    return seen


def parse_upload(content_base64, base_ds, dataset_id):
    """Validate the workbook and return (dataset, batches). Raises WorkbookErrors with row-level messages."""
    errors = []
    try:
        wb = load_workbook(io.BytesIO(base64.b64decode(content_base64, validate=True)), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise WorkbookErrors([{'sheet': None, 'row': None, 'column': None, 'message': 'File is not a readable .xlsx workbook'}])
    sheets = {name: _rows(wb, name, errors) for name in INPUT_SHEETS if name != 'Sources_Assumptions'}
    if errors:
        raise WorkbookErrors(errors)

    ds = copy.deepcopy(base_ds)
    ds['id'], ds['source'] = dataset_id, 'uploaded workbook'

    loc_ids = _unique(errors, 'Locations', sheets['Locations'], 'location_id')
    locations = {}
    for row, r in sheets['Locations']:
        lid = str(r['location_id'] or '').strip()
        typ = str(r['type'] or '').strip().lower()
        if typ not in ('hub', 'buyer'):
            errors.append({'sheet': 'Locations', 'row': row, 'column': 'type', 'message': 'Must be "hub" or "buyer"'})
        if typ == 'buyer' and str(r['channel'] or '') not in ('retail', 'horeca', 'wholesale', 'processor'):
            errors.append({'sheet': 'Locations', 'row': row, 'column': 'channel', 'message': 'Buyer channel must be retail, horeca, wholesale or processor'})
        lat = _num(errors, 'Locations', row, 'lat', r['lat'], 18.5, 19.6)
        lon = _num(errors, 'Locations', row, 'lon', r['lon'], 72.6, 73.5)
        if lid:
            locations[lid] = {'id': lid, 'name': str(r['name'] or lid), 'type': typ, 'channel': str(r['channel'] or ''),
                              'area': str(r['area'] or ''), 'lat': lat, 'lon': lon}
    hubs = {k for k, v in locations.items() if v['type'] == 'hub'}
    buyers_set = {k for k, v in locations.items() if v['type'] == 'buyer'}
    unknown_network = sorted((hubs | buyers_set) - set(base_ds['locations']))
    if unknown_network:
        errors.append({'sheet': 'Locations', 'row': None, 'column': 'location_id',
                       'message': f'New locations {unknown_network} have no road geometry in the built-in network; add them via backend/build_mumbai_network.py'})

    commodity_rows = sheets['Commodities']
    _unique(errors, 'Commodities', commodity_rows, 'commodity_id')
    fresh = {str(r['commodity_id']).strip(): (row, r) for row, r in sheets['Freshness_Parameters'] if r.get('commodity_id')}
    commodities = {}
    for row, r in commodity_rows:
        cid = str(r['commodity_id'] or '').strip()
        price = _num(errors, 'Commodities', row, 'reference_price_inr_kg', r['reference_price_inr_kg'], positive=True)
        if cid not in fresh:
            errors.append({'sheet': 'Freshness_Parameters', 'row': None, 'column': 'commodity_id', 'message': f'No freshness parameters for "{cid}"'})
            continue
        frow, f = fresh[cid]
        shelf = _num(errors, 'Freshness_Parameters', frow, 'shelf_life_hours_20c', f['shelf_life_hours_20c'], 4, 5000)
        q10 = _num(errors, 'Freshness_Parameters', frow, 'q10', f['q10'], 1.0, 5.0)
        hum = _num(errors, 'Freshness_Parameters', frow, 'humidity_sensitivity', f['humidity_sensitivity'], 0, 0.9)
        commodities[cid] = {'id': cid, 'name': str(r['name'] or cid), 'category': str(r['category'] or ''),
                            'reference_price_inr_kg': price, 'shelf_life_hours_20c': shelf, 'q10': q10, 'humidity_sensitivity': hum}

    offers = {str(r['buyer_id']).strip(): (row, r) for row, r in sheets['Buyer_Offers'] if r.get('buyer_id')}
    caps = {str(r['buyer_id']).strip(): (row, r) for row, r in sheets['Buyer_Capacity'] if r.get('buyer_id')}
    buyers = {}
    for bid in sorted(buyers_set):
        if bid not in offers:
            errors.append({'sheet': 'Buyer_Offers', 'row': None, 'column': 'buyer_id', 'message': f'Buyer "{bid}" has no offer row'})
            continue
        if bid not in caps:
            errors.append({'sheet': 'Buyer_Capacity', 'row': None, 'column': 'buyer_id', 'message': f'Buyer "{bid}" has no capacity row'})
            continue
        orow, o = offers[bid]
        crow, c = caps[bid]
        grades = str(o['accepted_grades'] or '').replace(' ', '')
        if not grades or any(g not in 'ABC' for g in grades.split(',')):
            errors.append({'sheet': 'Buyer_Offers', 'row': orow, 'column': 'accepted_grades', 'message': 'Use a comma list of A, B, C'})
        buyers[bid] = {'buyer_id': bid, 'channel': locations[bid]['channel'],
                       'price_factor': _num(errors, 'Buyer_Offers', orow, 'price_factor_vs_reference', o['price_factor_vs_reference'], 0.05, 5),
                       'min_rsl_fraction': _num(errors, 'Buyer_Offers', orow, 'min_rsl_fraction_of_shelf_life', o['min_rsl_fraction_of_shelf_life'], 0, 1),
                       'accepted_grades': grades,
                       'capacity_kg': _num(errors, 'Buyer_Capacity', crow, 'capacity_kg', c['capacity_kg'], 0, 1e7)}

    routes = {}
    for row, r in sheets['Road_Routes']:
        o, b = str(r['origin_id'] or '').strip(), str(r['buyer_id'] or '').strip()
        if o not in hubs or b not in buyers_set:
            errors.append({'sheet': 'Road_Routes', 'row': row, 'column': 'origin_id/buyer_id', 'message': f'{o} → {b} must join a hub to a buyer listed in Locations'})
            continue
        routes[f'{o}>{b}'] = {'distance_km': _num(errors, 'Road_Routes', row, 'distance_km', r['distance_km'], positive=True),
                              'freeflow_minutes': _num(errors, 'Road_Routes', row, 'freeflow_minutes', r['freeflow_minutes'], positive=True)}

    costs = dict(ds['costs'])
    for row, r in sheets['Operating_Costs']:
        item = str(r['cost_item'] or '').strip()
        if item not in costs:
            errors.append({'sheet': 'Operating_Costs', 'row': row, 'column': 'cost_item', 'message': f'Unknown cost item "{item}"'})
            continue
        costs[item] = _num(errors, 'Operating_Costs', row, 'value', r['value'], 0, 1e7)
    settings = dict(ds['settings'])
    for row, r in sheets['Scenario_Settings']:
        key = str(r['setting'] or '').strip()
        if key not in SETTING_BOUNDS:
            errors.append({'sheet': 'Scenario_Settings', 'row': row, 'column': 'setting', 'message': f'Unknown setting "{key}"'})
            continue
        lo, hi = SETTING_BOUNDS[key]
        settings[key] = _num(errors, 'Scenario_Settings', row, 'value', r['value'], lo, hi)

    batch_ids = _unique(errors, 'Batch_Inventory', sheets['Batch_Inventory'], 'batch_id')
    batches = []
    for row, r in sheets['Batch_Inventory']:
        cid, o, b = (str(r[k] or '').strip() for k in ['commodity_id', 'origin_id', 'committed_buyer_id'])
        if cid not in commodities:
            errors.append({'sheet': 'Batch_Inventory', 'row': row, 'column': 'commodity_id', 'message': f'Unknown commodity "{cid}"'})
        if o not in hubs:
            errors.append({'sheet': 'Batch_Inventory', 'row': row, 'column': 'origin_id', 'message': f'"{o}" is not a hub in Locations'})
        if b not in buyers_set:
            errors.append({'sheet': 'Batch_Inventory', 'row': row, 'column': 'committed_buyer_id', 'message': f'"{b}" is not a buyer in Locations'})
        elif o in hubs and f'{o}>{b}' not in routes:
            errors.append({'sheet': 'Batch_Inventory', 'row': row, 'column': 'committed_buyer_id', 'message': f'No road route {o} → {b} in Road_Routes'})
        qty = _num(errors, 'Batch_Inventory', row, 'quantity_kg', r['quantity_kg'], positive=True, hi=10000)
        grade = str(r['quality_grade'] or 'A').strip().upper()
        if grade not in ('A', 'B', 'C'):
            errors.append({'sheet': 'Batch_Inventory', 'row': row, 'column': 'quality_grade', 'message': 'Must be A, B or C'})
        batches.append({'batch_id': str(r['batch_id']).strip(), 'commodity_id': cid, 'origin_id': o, 'buyer_id': b,
                        'quantity_kg': qty, 'quality_grade': grade, 'disruptions': []})
    by_id = {b['batch_id']: b for b in batches}
    for row, r in sheets['Disruption_Events']:
        bid, typ, sev = (str(r[k] or '').strip() for k in ['batch_id', 'disruption_type', 'severity'])
        if bid not in batch_ids:
            errors.append({'sheet': 'Disruption_Events', 'row': row, 'column': 'batch_id', 'message': f'Unknown batch "{bid}"'})
            continue
        if typ not in DISRUPTIONS:
            errors.append({'sheet': 'Disruption_Events', 'row': row, 'column': 'disruption_type', 'message': f'Unknown disruption "{typ}"'})
            continue
        if sev.lower() not in SEVERITIES:
            errors.append({'sheet': 'Disruption_Events', 'row': row, 'column': 'severity', 'message': 'Must be low, medium or high'})
            continue
        if any(d['type'] == typ for d in by_id[bid]['disruptions']):
            errors.append({'sheet': 'Disruption_Events', 'row': row, 'column': 'disruption_type', 'message': f'"{typ}" listed twice for {bid}'})
            continue
        by_id[bid]['disruptions'].append({'type': typ, 'severity': sev.lower()})
    if not batches:
        errors.append({'sheet': 'Batch_Inventory', 'row': None, 'column': None, 'message': 'Add at least one batch'})
    if errors:
        raise WorkbookErrors(errors)

    ds.update(locations={k: v for k, v in ds['locations'].items() if k in locations}, commodities=commodities, buyers=buyers,
              routes=routes, costs=costs, settings=settings)
    for k, v in locations.items():
        ds['locations'][k] = {**ds['locations'].get(k, {}), **v}
    return ds, batches


# ----------------------------------------------------------------------------------- exports

STRATEGY_COLUMNS = ['strategy', 'model', 'label', 'status', 'buyer_id', 'buyer_name', 'diverted', 'lead_time_hours', 'arrival_rsl_hours',
                    'minimum_rsl_hours', 'accepted', 'input_kg', 'sellable_kg', 'spoiled_kg', 'offer_inr_kg', 'price_inr_kg', 'revenue_inr',
                    'procurement_inr', 'transport_inr', 'handling_inr', 'rescue_inr', 'agent_tech_inr', 'gross_margin_inr',
                    'operating_profit_inr', 'distance_km']


def shipment_export(items, title):
    """items: list of {'id','request','settings','view','state','dataset'} for one or many shipments."""
    wb = Workbook()
    wb.remove(wb.active)
    summary = []
    comparison, decisions, routes, fresh, pricing, ledger, actions, inputs = [], [], [], [], [], [], [], []
    for it in items:
        s = it['view']['strategies']
        req = it['request']
        a, d = s['A'], s['D']
        summary.append({'run_id': it['id'], 'batch_id': it.get('batch_id', ''), 'commodity': req['commodity_id'], 'origin': req['origin_id'],
                        'committed_buyer': req['buyer_id'], 'quantity_kg': req['quantity_kg'], 'seed': it['settings']['seed'],
                        'dataset': it['state']['shipment']['dataset_id'], 'disruptions': ', '.join(f"{x['type']}:{x['severity']}" for x in req['disruptions']) or 'none',
                        'agent_mode': it['state'].get('agent_mode'), 'decision_status': it['view']['decision_status'],
                        'A_profit_inr': a['operating_profit_inr'], 'D_profit_inr': d['operating_profit_inr'],
                        'D_minus_A_inr': d['operating_profit_inr'] - a['operating_profit_inr'],
                        'D_minus_C_inr': d['operating_profit_inr'] - s['C']['operating_profit_inr'],
                        'explanation': ' '.join(it['view']['narrative'])})
        for k in 'ABCD':
            row = {'run_id': it['id'], 'batch_id': it.get('batch_id', ''), **{c: s[k].get(c) for c in STRATEGY_COLUMNS}}
            comparison.append(row)
            decisions.append({'run_id': it['id'], 'strategy': k, 'buyer_id': s[k]['buyer_id'], 'buyer_name': s[k]['buyer_name'],
                              'accepted': s[k]['accepted'], 'reasons': '; '.join(s[k]['rejection_reasons']) or 'accepted'})
            routes.append({'run_id': it['id'], 'strategy': k, 'route': s[k]['route_key'], 'diversion_leg': ' → '.join(s[k]['leg'] or []) or '',
                           'distance_km': s[k]['distance_km'], 'lead_time_hours': s[k]['lead_time_hours']})
            for p in s[k]['timeline']:
                fresh.append({'run_id': it['id'], 'strategy': k, 'hours_since_hub': p['hours'], 'remaining_shelf_life_hours': p['rsl_hours'], 'event': p['label']})
            pricing.append({'run_id': it['id'], 'strategy': k, 'buyer_id': s[k]['buyer_id'], 'offer_inr_kg': s[k]['offer_inr_kg'],
                            'realized_inr_kg': s[k]['price_inr_kg'], 'sellable_kg': s[k]['sellable_kg'], 'revenue_inr': s[k]['revenue_inr']})
            for acct in ['revenue_inr', 'procurement_inr', 'transport_inr', 'handling_inr', 'rescue_inr', 'agent_tech_inr']:
                sign = 1 if acct == 'revenue_inr' else -1
                ledger.append({'run_id': it['id'], 'strategy': k, 'account': acct, 'amount_inr': sign * s[k][acct]})
        for t in it['state']['trace']:
            actions.append({'run_id': it['id'], 'sequence': t['sequence'], 'timestamp': t['timestamp'], 'tool': t['tool'],
                            'inputs': str(t['inputs'])[:30000], 'result': str(t['result'])[:30000]})
        for k, v in {**req, **{f'setting.{a}': b for a, b in it['settings'].items()}}.items():
            inputs.append({'run_id': it['id'], 'input': k, 'value': str(v)})

    head = wb.create_sheet('Scenario_Summary')
    head.append([title])
    head.append([f'Exported {datetime.now(timezone.utc).isoformat(timespec="seconds")}. All values are SIMULATED RESULT from synthetic inputs.'])
    head.append([])
    cols = list(summary[0]) if summary else ['run_id']
    head.append(cols)
    for r in summary:
        head.append([r[c] for c in cols])
    style(head, header_row=4)
    head['A1'].font = Font(bold=True, size=14, color='1D4433')
    for name, rows in [('A_B_C_D_Comparison', comparison), ('Batch_Outcomes', comparison), ('Buyer_Decisions', decisions),
                       ('Route_ETA_Changes', routes), ('Freshness_Quality', fresh), ('Realized_Pricing', pricing),
                       ('Financial_Ledger', ledger), ('Agent_Action_Log', actions), ('Scenario_Inputs', inputs)]:
        add_sheet(wb, name, list(rows[0]) if rows else ['run_id'], rows)
    add_sheet(wb, 'Sources_Assumptions', INPUT_SHEETS['Sources_Assumptions'][0], sources_rows())
    return to_bytes(wb)


def fleet_export(run_id, result):
    wb = Workbook()
    wb.remove(wb.active)
    letters = {'M0': 'A', 'M1': 'B', 'M2': 'C', 'M3': 'D'}
    summary = [{'strategy': letters[m], 'model': m, **{k: v for k, v in s.items() if not isinstance(v, list)}} for m, s in result['summary'].items()]
    add_sheet(wb, 'Scenario_Summary', list(summary[0]), summary)
    metrics = []
    for rep in result['runs']:
        for m in rep['metrics']:
            metrics.append({'run_id': run_id, 'seed': rep['seed'], 'dataset_sha256': rep['dataset_sha256'], 'strategy': letters[m['model']],
                            **{k: v for k, v in m.items() if not isinstance(v, list)}})
    add_sheet(wb, 'A_B_C_D_Comparison', list(metrics[0]), metrics)
    ledger = [{'seed': rep['seed'], 'strategy': letters[m['model']], **line} for rep in result['runs'] for m in rep['metrics'] for line in m.get('financial_waterfall', [])]
    add_sheet(wb, 'Financial_Ledger', list(ledger[0]) if ledger else ['seed'], ledger)
    add_sheet(wb, 'Scenario_Inputs', ['input', 'value'], [{'input': k, 'value': v} for k, v in result['config'].items()])
    add_sheet(wb, 'Limitations', ['limitation'], [{'limitation': x} for x in result['limitations']])
    add_sheet(wb, 'Sources_Assumptions', INPUT_SHEETS['Sources_Assumptions'][0], sources_rows())
    return to_bytes(wb)
