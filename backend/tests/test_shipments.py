import base64
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app

HEAT_AND_TRAFFIC = [{'type': 'heavy_traffic', 'severity': 'high'}, {'type': 'temperature_rise', 'severity': 'high'}]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv('FRESHFLOW_OPERATOR_KEY', raising=False)
    monkeypatch.delenv('OPENAI_API_KEY', raising=False)
    monkeypatch.setenv('FRESHFLOW_ENV', 'development')
    return TestClient(create_app(tmp_path / 'ship.sqlite'))


def post(client, **body):
    r = client.post('/api/v1/shipments', json=body)
    assert r.status_code == 200, r.text
    return r.json()


def tools(payload):
    return [t['tool'] for t in payload['state']['trace']]


def test_network_has_15_locations_and_osrm_routes(client):
    net = client.get('/api/v1/network').json()
    assert len(net['hubs']) + len(net['buyers']) == 15
    assert len(net['commodities']) == 18
    route = net['routes']['H1>B02']
    assert route['distance_km'] > 0 and len(route['geometry']) > 2
    assert 'OSRM' in net['source'] and 'Not live' in net['traffic']


def test_no_disruption_needs_no_intervention_and_no_fake_approval(client):
    p = post(client)
    assert p['decision_status'] == 'no_intervention_needed'
    assert 'request_human_approval' not in tools(p)
    assert tools(p)[-1] == 'no_intervention_required'
    assert p['strategies']['A']['buyer_id'] == p['strategies']['D']['buyer_id'] == 'B02'


def test_disruption_triggers_real_tool_calls_and_approval_executes(client):
    p = post(client, disruptions=HEAT_AND_TRAFFIC)
    s = p['strategies']
    assert not s['A']['accepted']
    assert p['decision_status'] == 'proposed_awaiting_approval'
    assert tools(p)[2:] == ['get_batch_state', 'predict_remaining_shelf_life', 'get_buyer_options', 'check_buyer_capacity',
                            'calculate_route_eta', 'optimize_destination', 'validate_constraints', 'request_human_approval']
    target = p['state']['proposal']['id']
    assert target != 'B02'
    r = client.post(f"/api/v1/shipments/{p['id']}/approval", json={'approve': True, 'expected_version': p['state']['version']})
    assert r.status_code == 200
    done = r.json()
    assert done['decision_status'] == 'executed'
    assert done['state']['destination'] == target
    assert 'execute_destination_change' in tools(done) and 'verify_execution' in tools(done)
    stale = client.post(f"/api/v1/shipments/{p['id']}/approval", json={'approve': True, 'expected_version': p['state']['version']})
    assert stale.status_code == 409


def test_reject_keeps_plan(client):
    p = post(client, disruptions=HEAT_AND_TRAFFIC)
    r = client.post(f"/api/v1/shipments/{p['id']}/approval", json={'approve': False, 'expected_version': p['state']['version']}).json()
    assert r['decision_status'] == 'rejected_kept_plan'
    assert r['state']['destination'] == 'B02'
    assert r['strategies']['D']['buyer_id'] == 'B02'


def test_financials_reconcile_for_every_strategy(client):
    p = post(client, disruptions=HEAT_AND_TRAFFIC + [{'type': 'product_damage', 'severity': 'medium'}])
    for s in p['strategies'].values():
        costs = s['procurement_inr'] + s['transport_inr'] + s['handling_inr'] + s['rescue_inr'] + s['agent_tech_inr']
        assert abs(s['revenue_inr'] - costs - s['operating_profit_inr']) < 1e-6
        assert abs(s['sellable_kg'] + s['spoiled_kg'] - s['input_kg']) < 1e-6
        assert s['revenue_inr'] == pytest.approx(s['sellable_kg'] * s['price_inr_kg'])


def test_unrescuable_shipment_escalates_instead_of_forcing_a_diversion(client):
    p = post(client, commodity_id='spinach', buyer_id='B10', settings={'initial_shelf_life_multiplier': 0.3},
             disruptions=[{'type': 'refrigeration_failure', 'severity': 'high'}, {'type': 'road_closure', 'severity': 'high'}])
    assert p['decision_status'] == 'no_feasible_alternative'
    assert 'escalate' in tools(p)
    assert not p['strategies']['D']['accepted']


def test_reproducible(client):
    a = post(client, disruptions=HEAT_AND_TRAFFIC, settings={'seed': 99})
    b = post(client, disruptions=HEAT_AND_TRAFFIC, settings={'seed': 99})
    assert a['strategies'] == b['strategies']


def test_live_mode_without_key_fails_loudly(client):
    r = client.post('/api/v1/shipments', json={'mode': 'live'})
    assert r.status_code == 409 and 'not configured' in r.json()['detail']


def test_invalid_inputs_rejected(client):
    assert client.post('/api/v1/shipments', json={'commodity_id': 'durian'}).status_code == 422
    assert client.post('/api/v1/shipments', json={'origin_id': 'B02'}).status_code == 422
    assert client.post('/api/v1/shipments', json={'settings': {'traffic_multiplier': 99}}).status_code == 422
    assert client.post('/api/v1/shipments', json={'disruptions': [{'type': 'heavy_traffic'}, {'type': 'heavy_traffic'}]}).status_code == 422


def b64(content):
    return base64.b64encode(content).decode()


def test_excel_template_round_trip_runs_uploaded_batches(client):
    template = client.get('/api/v1/excel/template')
    assert template.status_code == 200
    wb = load_workbook(io.BytesIO(template.content))
    for sheet in ['Locations', 'Road_Routes', 'Commodities', 'Batch_Inventory', 'Buyer_Offers', 'Buyer_Capacity',
                  'Freshness_Parameters', 'Disruption_Events', 'Operating_Costs', 'Scenario_Settings', 'Sources_Assumptions']:
        assert sheet in wb.sheetnames
    # modify: double the tomato price, so the uploaded data must actually be used
    ws = wb['Commodities']
    for row in ws.iter_rows(min_row=2):
        if row[0].value == 'tomato':
            row[3].value = 60
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post('/api/v1/excel/upload', json={'filename': 'edited.xlsx', 'content_base64': b64(buf.getvalue())})
    assert r.status_code == 200, r.text
    run = r.json()
    assert run['totals']['batches'] == 3
    tomato = next(s for s in run['shipments'] if s['batch_id'] == 'B-001')
    shipment = client.get(f"/api/v1/shipments/{tomato['shipment_id']}").json()
    assert shipment['commodity']['reference_price_inr_kg'] == 60
    assert shipment['dataset_id'] == run['dataset_id'] != 'default'
    export = client.get(f"/api/v1/workbook-runs/{run['id']}/export.xlsx")
    out = load_workbook(io.BytesIO(export.content))
    assert {'Scenario_Summary', 'A_B_C_D_Comparison', 'Agent_Action_Log', 'Financial_Ledger'} <= set(out.sheetnames)
    assert out['A_B_C_D_Comparison'].max_row == 1 + 3 * 4


def test_bad_workbook_reports_row_level_errors(client):
    wb = load_workbook(io.BytesIO(client.get('/api/v1/excel/template').content))
    inv = wb['Batch_Inventory']
    inv.append(['B-001', 'durian', 'H9', 'B02', -5, 'Z'])  # duplicate id + 4 bad values
    del wb['Buyer_Capacity']
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post('/api/v1/excel/upload', json={'filename': 'bad.xlsx', 'content_base64': b64(buf.getvalue())})
    assert r.status_code == 422
    errors = r.json()['detail']['errors']
    assert any(e['sheet'] == 'Buyer_Capacity' and 'missing' in e['message'].lower() for e in errors)
    r2 = client.post('/api/v1/excel/upload', json={'filename': 'x.xlsx', 'content_base64': b64(b'not a workbook')})
    assert r2.status_code == 422


def test_bad_rows_named_precisely(client):
    wb = load_workbook(io.BytesIO(client.get('/api/v1/excel/template').content))
    wb['Batch_Inventory'].append(['B-001', 'durian', 'H9', 'B02', -5, 'Z'])
    buf = io.BytesIO()
    wb.save(buf)
    errors = client.post('/api/v1/excel/upload', json={'filename': 'bad.xlsx', 'content_base64': b64(buf.getvalue())}).json()['detail']['errors']
    messages = {(e['sheet'], e['row'], e['column']): e['message'] for e in errors}
    assert ('Batch_Inventory', 5, 'batch_id') in messages and 'Duplicate' in messages[('Batch_Inventory', 5, 'batch_id')]
    assert ('Batch_Inventory', 5, 'commodity_id') in messages
    assert ('Batch_Inventory', 5, 'quantity_kg') in messages
    assert ('Batch_Inventory', 5, 'quality_grade') in messages


def test_shipment_and_comparison_exports(client):
    a = post(client, disruptions=HEAT_AND_TRAFFIC)
    b = post(client)
    one = load_workbook(io.BytesIO(client.get(f"/api/v1/shipments/{a['id']}/export.xlsx").content))
    assert one['Agent_Action_Log'].max_row - 1 == len(a['state']['trace'])
    both = client.post('/api/v1/exports/comparison', json={'shipment_ids': [a['id'], b['id']]})
    wb = load_workbook(io.BytesIO(both.content))
    ids = {wb['Scenario_Summary'].cell(row=r, column=1).value for r in range(5, wb['Scenario_Summary'].max_row + 1)}
    assert ids == {a['id'], b['id']}


def test_number_formatting_does_not_change_results(client):
    a = post(client, quantity_kg=800, settings={'seed': 7})
    b = post(client, quantity_kg=800.0, settings={'seed': 7.0})
    assert a['strategies'] == b['strategies'] and a['plan_buyer_id'] == b['plan_buyer_id']


def test_live_provider_failure_is_reported_not_disguised(client, monkeypatch):
    import httpx
    monkeypatch.setenv('OPENAI_API_KEY', 'test-key')
    monkeypatch.setenv('OPENAI_MODEL', 'test-model')
    from app import agent
    provider = httpx.MockTransport(lambda request: httpx.Response(429, json={'error': {'message': 'You have no credits remaining.'}}))
    original = agent.live_loop
    monkeypatch.setattr(agent, 'live_loop', lambda inv, transport=None: original(inv, provider))
    p = post(client, mode='live', disruptions=HEAT_AND_TRAFFIC)
    assert p['decision_status'] == 'agent_failed'
    failure = next(t for t in p['state']['trace'] if t['tool'] == 'investigation_failed')
    assert 'no credits' in failure['result']['message'] and 'test-key' not in str(p)
    assert p['strategies']['D']['buyer_id'] == p['plan_buyer_id']  # nothing executed
